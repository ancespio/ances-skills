#!/usr/bin/env python3
"""Validate pptx-workshop scenario contracts and hard gates."""

from __future__ import annotations

import argparse
import binascii
import hashlib
import json
import math
import re
import struct
import tempfile
import zipfile
import zlib
from datetime import datetime, timezone
from pathlib import Path
from typing import Any


SCENARIOS = {
    "new-deck",
    "template-fill",
    "reconstruction",
    "reference-deck",
    "existing-deck",
}
PHASES = {"plan", "preview", "final"}
SHA_RE = re.compile(r"^[0-9a-f]{64}$")
SLIDE_RE = re.compile(r"^P\d{2,3}$")
BLOCK_RE = re.compile(r"^B\d{2,3}$")
SKILL_ROOT = Path(__file__).resolve().parent.parent
STYLE_LIBRARY_CATALOG = SKILL_ROOT / "assets" / "style-library" / "catalog.json"
STYLE_LIBRARY_VERSION = "pptx-workshop.style-library.v1"
DOCS = {
    "brief": "plans/brief.json",
    "outline": "plans/outline.json",
    "slide_plan": "plans/slide_plan.json",
    "evidence_plan": "plans/evidence_plan.json",
    "asset_plan": "plans/asset_plan.json",
    "gorden_component": "plans/gorden_component.json",
    "style_reference_plan": "plans/style_reference_plan.json",
    "template_analysis": "plans/template_analysis.json",
    "adaptation_report": "plans/adaptation_report.json",
    "fill_plan": "plans/fill_plan.json",
    "reconstruction_plan": "plans/reconstruction_plan.json",
    "reference_profile": "plans/reference_profile.json",
    "change_plan": "plans/change_plan.json",
    "scene1_handoff": "plans/scene1_handoff.json",
    "build_plan": "plans/build_plan.json",
    "artifacts": "manifests/pptx_artifacts.json",
    "report": "manifests/editability_report.json",
}
SCHEMA_VERSIONS = {
    "brief": "pptx-workshop.brief.v2",
    "outline": "pptx-workshop.outline.v2",
    "slide_plan": "pptx-workshop.slide-plan.v3",
    "evidence_plan": "pptx-workshop.evidence-plan.v1",
    "asset_plan": "pptx-workshop.asset-plan.v2",
    "gorden_component": "pptx-workshop.gorden-component.v1",
    "style_reference_plan": "pptx-workshop.style-reference-plan.v1",
    "template_analysis": "pptx-workshop.template-analysis.v1",
    "adaptation_report": "pptx-workshop.adaptation-report.v1",
    "fill_plan": "pptx-workshop.fill-plan.v1",
    "reconstruction_plan": "pptx-workshop.reconstruction-plan.v2",
    "reference_profile": "pptx-workshop.reference-profile.v1",
    "change_plan": "pptx-workshop.change-plan.v1",
    "scene1_handoff": "pptx-workshop.scene1-handoff.v2",
    "build_plan": "pptx-workshop.build-plan.v2",
    "artifacts": "pptx-workshop.pptx-artifacts.v1",
    "report": "pptx-workshop.editability-report.v3",
}
PLAN_DOCS = {
    "new-deck": ["brief", "evidence_plan", "outline", "slide_plan", "asset_plan", "gorden_component"],
    "template-fill": ["brief", "template_analysis", "adaptation_report", "fill_plan"],
    "reconstruction": ["brief", "reconstruction_plan", "asset_plan", "gorden_component"],
    "reference-deck": [
        "brief", "reference_profile", "evidence_plan", "outline", "slide_plan", "asset_plan"
    ],
    "existing-deck": ["brief", "change_plan"],
}
CONFIRMED_STATUS = "confirmed"
GORDEN_GENERATION_MANIFEST_VERSION = "pptx-workshop.gorden-generation-manifest.v3"


class ContractError(ValueError):
    """Raised when one or more hard gates fail."""


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def json_sha256(value: Any) -> str:
    payload = json.dumps(
        value,
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
    ).encode("utf-8")
    return hashlib.sha256(payload).hexdigest()


def parse_timestamp(value: Any, context: str) -> datetime:
    if not isinstance(value, str) or not value.strip():
        raise ContractError(f"{context} must be a non-empty ISO 8601 timestamp")
    try:
        parsed = datetime.fromisoformat(value.replace("Z", "+00:00"))
    except ValueError as exc:
        raise ContractError(f"{context} must be an ISO 8601 timestamp") from exc
    if parsed.tzinfo is None:
        raise ContractError(f"{context} must include a timezone")
    return parsed.astimezone(timezone.utc)


def png_dimensions(path: Path) -> tuple[int, int]:
    data = path.read_bytes()
    if not data.startswith(b"\x89PNG\r\n\x1a\n"):
        raise ContractError(f"not a PNG file: {path}")
    offset = 8
    width = height = 0
    seen_idat = False
    seen_iend = False
    while offset + 12 <= len(data):
        length = struct.unpack(">I", data[offset : offset + 4])[0]
        chunk_type = data[offset + 4 : offset + 8]
        end = offset + 12 + length
        if end > len(data):
            raise ContractError(f"truncated PNG chunk: {path}")
        payload = data[offset + 8 : offset + 8 + length]
        expected_crc = struct.unpack(">I", data[offset + 8 + length : end])[0]
        actual_crc = binascii.crc32(chunk_type + payload) & 0xFFFFFFFF
        if actual_crc != expected_crc:
            raise ContractError(f"invalid PNG chunk CRC: {path}")
        if chunk_type == b"IHDR":
            if offset != 8 or length != 13:
                raise ContractError(f"invalid PNG IHDR: {path}")
            width, height = struct.unpack(">II", payload[:8])
        elif chunk_type == b"IDAT":
            seen_idat = True
        elif chunk_type == b"IEND":
            seen_iend = True
            break
        offset = end
    if width < 1 or height < 1 or not seen_idat or not seen_iend:
        raise ContractError(f"incomplete PNG file: {path}")
    return width, height


def load_json(path: Path) -> dict[str, Any]:
    try:
        value = json.loads(path.read_text(encoding="utf-8"))
    except FileNotFoundError as exc:
        raise ContractError(f"missing file: {path}") from exc
    except (OSError, UnicodeError, json.JSONDecodeError) as exc:
        raise ContractError(f"invalid UTF-8 JSON: {path}: {exc}") from exc
    if not isinstance(value, dict):
        raise ContractError(f"top-level JSON must be an object: {path}")
    return value


def require(obj: dict[str, Any], fields: list[str], context: str) -> None:
    missing = [field for field in fields if field not in obj]
    if missing:
        raise ContractError(f"{context} missing fields: {', '.join(missing)}")


def resolve_inside(root: Path, relative: str, context: str) -> Path:
    path = (root / relative).resolve()
    try:
        path.relative_to(root)
    except ValueError as exc:
        raise ContractError(f"{context} must stay inside run directory: {relative}") from exc
    return path


def load_doc(root: Path, key: str, run_id: str) -> dict[str, Any]:
    doc = load_json(root / DOCS[key])
    expected = SCHEMA_VERSIONS[key]
    if doc.get("schema_version") != expected:
        raise ContractError(f"{DOCS[key]}.schema_version must be {expected}")
    if doc.get("run_id") != run_id:
        raise ContractError(f"{DOCS[key]}.run_id does not match project.json")
    if doc.get("status") == "confirmed_by_delegation":
        raise ContractError(f"{DOCS[key]} may not use confirmed_by_delegation")
    return doc


def validate_style_reference_plan(root: Path, plan: dict[str, Any]) -> None:
    require(
        plan,
        [
            "catalog_schema_version",
            "catalog_sha256",
            "candidates",
            "selected_ids",
            "usage_mode",
            "user_direction",
        ],
        "style_reference_plan.json",
    )
    if not STYLE_LIBRARY_CATALOG.is_file():
        raise ContractError("bundled style-library catalog is missing")
    if (
        plan["catalog_schema_version"] != STYLE_LIBRARY_VERSION
        or plan["catalog_sha256"] != sha256(STYLE_LIBRARY_CATALOG)
    ):
        raise ContractError("style_reference_plan.json is not bound to the current style-library catalog")
    catalog = load_json(STYLE_LIBRARY_CATALOG)
    if catalog.get("schema_version") != STYLE_LIBRARY_VERSION:
        raise ContractError("bundled style-library catalog schema is invalid")
    catalog_entries = {
        item.get("id"): item
        for item in catalog.get("entries", [])
        if isinstance(item, dict) and isinstance(item.get("id"), str)
    }

    candidates = plan["candidates"]
    if not isinstance(candidates, list) or not 3 <= len(candidates) <= 9:
        raise ContractError("style_reference_plan.candidates must contain 3 to 9 shown choices")
    candidate_ids: list[str] = []
    candidate_sources: dict[str, str] = {}
    for index, candidate in enumerate(candidates):
        context = f"style_reference_plan.candidates[{index}]"
        if not isinstance(candidate, dict):
            raise ContractError(f"{context} must be an object")
        require(candidate, ["id", "source_type", "style_label", "fit_reason", "limitations"], context)
        if any(not isinstance(candidate[field], str) or not candidate[field].strip() for field in ("id", "style_label", "fit_reason")):
            raise ContractError(f"{context} id, style_label and fit_reason must not be blank")
        if not isinstance(candidate["limitations"], list) or any(
            not isinstance(item, str) or not item.strip() for item in candidate["limitations"]
        ):
            raise ContractError(f"{context}.limitations must be an array of non-blank strings")
        source_type = candidate["source_type"]
        if source_type == "bundled":
            require(candidate, ["catalog_entry_id"], context)
            catalog_id = candidate["catalog_entry_id"]
            if candidate["id"] != catalog_id or catalog_id not in catalog_entries:
                raise ContractError(f"{context} names an unknown bundled catalog entry")
            if "new-deck-style-reference" not in catalog_entries[catalog_id].get("scene_eligibility", []):
                raise ContractError(f"{context} catalog entry is not eligible for scene 1")
        elif source_type == "user-provided":
            require(candidate, ["source_path", "source_sha256"], context)
            source_path = resolve_inside(root, candidate["source_path"], f"{context}.source_path")
            if not source_path.is_file() or candidate["source_sha256"] != sha256(source_path):
                raise ContractError(f"{context} user-provided reference is missing or stale")
        elif source_type != "custom":
            raise ContractError(f"{context}.source_type is invalid")
        candidate_ids.append(candidate["id"])
        candidate_sources[candidate["id"]] = source_type
    if len(candidate_ids) != len(set(candidate_ids)):
        raise ContractError("style_reference_plan.candidates contains duplicate ids")

    selected_ids = plan["selected_ids"]
    if (
        not isinstance(selected_ids, list)
        or not 1 <= len(selected_ids) <= 3
        or len(selected_ids) != len(set(selected_ids))
        or any(item not in candidate_ids for item in selected_ids)
    ):
        raise ContractError("style_reference_plan.selected_ids must select 1 to 3 shown candidates")
    usage_mode = plan["usage_mode"]
    if usage_mode not in {"direct", "blend", "inspired", "custom"}:
        raise ContractError("style_reference_plan.usage_mode is invalid")
    if usage_mode == "direct" and len(selected_ids) != 1:
        raise ContractError("direct style usage must select exactly one reference")
    if usage_mode == "blend" and len(selected_ids) < 2:
        raise ContractError("blended style usage must select at least two references")
    if usage_mode == "custom" and any(candidate_sources[item] != "custom" for item in selected_ids):
        raise ContractError("custom style usage may select only custom candidates")
    if not isinstance(plan["user_direction"], str) or not plan["user_direction"].strip():
        raise ContractError("style_reference_plan.user_direction must not be blank")


def validate_project_file(root: Path) -> dict[str, Any]:
    project = load_json(root / "project.json")
    require(
        project,
        ["schema_version", "run_id", "scenario", "phase", "source_files", "new_page_count"],
        "project.json",
    )
    if project["schema_version"] != "pptx-workshop.project.v1":
        raise ContractError("project.json.schema_version is invalid")
    if project["scenario"] not in SCENARIOS:
        raise ContractError("project.json.scenario is invalid")
    if not isinstance(project["new_page_count"], int) or project["new_page_count"] < 0:
        raise ContractError("project.json.new_page_count must be a non-negative integer")
    sources = project["source_files"]
    if not isinstance(sources, list):
        raise ContractError("project.json.source_files must be an array")
    for index, source in enumerate(sources):
        if not isinstance(source, dict):
            raise ContractError(f"project.json.source_files[{index}] must be an object")
        require(source, ["path", "sha256", "role"], f"project.json.source_files[{index}]")
        path = Path(source["path"]).expanduser().resolve()
        if not path.is_file():
            raise ContractError(f"source file does not exist: {path}")
        if source["sha256"] != sha256(path):
            raise ContractError(f"source hash mismatch: {path}")
    return project


def slide_ids(doc: dict[str, Any], context: str) -> list[str]:
    slides = doc.get("slides")
    if not isinstance(slides, list) or not slides:
        raise ContractError(f"{context}.slides must be a non-empty array")
    ids: list[str] = []
    for index, slide in enumerate(slides):
        if not isinstance(slide, dict) or not SLIDE_RE.fullmatch(str(slide.get("id", ""))):
            raise ContractError(f"{context}.slides[{index}].id must match PNN or PNNN")
        ids.append(slide["id"])
    if len(ids) != len(set(ids)):
        raise ContractError(f"{context}.slides contains duplicate ids")
    return ids


def validate_fraction_bbox(value: Any, context: str) -> None:
    if (
        not isinstance(value, list)
        or len(value) != 4
        or any(isinstance(item, bool) or not isinstance(item, (int, float)) for item in value)
    ):
        raise ContractError(f"{context} must be [x,y,w,h] numeric fractions")
    x, y, width, height = (float(item) for item in value)
    if (
        x < 0
        or y < 0
        or width <= 0
        or height <= 0
        or x > 1
        or y > 1
        or x + width > 1 + 1e-9
        or y + height > 1 + 1e-9
    ):
        raise ContractError(f"{context} must stay inside the normalized slide canvas")


def validate_slide_plan(doc: dict[str, Any]) -> None:
    slide_ids(doc, "slide_plan.json")
    for slide_index, slide in enumerate(doc["slides"]):
        context = f"slide_plan.json.slides[{slide_index}]"
        require(
            slide,
            [
                "text_locked",
                "layout_locked",
                "content_blocks",
                "layout",
                "asset_ids",
                "native_objects",
                "qa_assertions",
            ],
            context,
        )
        if slide["text_locked"] is not True or slide["layout_locked"] is not True:
            raise ContractError(f"{context} must lock complete text and approximate layout")
        blocks = slide["content_blocks"]
        if not isinstance(blocks, list) or not blocks:
            raise ContractError(f"{context}.content_blocks must be a non-empty array")
        block_ids: list[str] = []
        reading_orders: list[int] = []
        for block_index, block in enumerate(blocks):
            block_context = f"{context}.content_blocks[{block_index}]"
            if not isinstance(block, dict):
                raise ContractError(f"{block_context} must be an object")
            require(block, ["id", "type", "text", "bbox", "reading_order", "align"], block_context)
            if not BLOCK_RE.fullmatch(str(block["id"])):
                raise ContractError(f"{block_context}.id must match BNN or BNNN")
            if any(not isinstance(block[field], str) or not block[field].strip() for field in ("type", "text", "align")):
                raise ContractError(f"{block_context} must contain final text, type and alignment")
            validate_fraction_bbox(block["bbox"], f"{block_context}.bbox")
            order = block["reading_order"]
            if isinstance(order, bool) or not isinstance(order, int) or order < 1:
                raise ContractError(f"{block_context}.reading_order must be a positive integer")
            block_ids.append(block["id"])
            reading_orders.append(order)
        if len(block_ids) != len(set(block_ids)):
            raise ContractError(f"{context}.content_blocks contains duplicate ids")
        if sorted(reading_orders) != list(range(1, len(blocks) + 1)):
            raise ContractError(f"{context}.content_blocks reading_order must be consecutive and unique")

        layout = slide["layout"]
        if not isinstance(layout, dict):
            raise ContractError(f"{context}.layout must be an object")
        require(layout, ["pattern", "reading_order", "decorative_safe_zones"], f"{context}.layout")
        if not isinstance(layout["pattern"], str) or not layout["pattern"].strip():
            raise ContractError(f"{context}.layout.pattern must not be blank")
        expected_order = [
            block["id"] for block in sorted(blocks, key=lambda item: item["reading_order"])
        ]
        if layout["reading_order"] != expected_order:
            raise ContractError(f"{context}.layout.reading_order must match the ordered content block ids")
        safe_zones = layout["decorative_safe_zones"]
        if not isinstance(safe_zones, list):
            raise ContractError(f"{context}.layout.decorative_safe_zones must be an array")
        zone_ids: list[str] = []
        for zone_index, zone in enumerate(safe_zones):
            zone_context = f"{context}.layout.decorative_safe_zones[{zone_index}]"
            if not isinstance(zone, dict):
                raise ContractError(f"{zone_context} must be an object")
            require(zone, ["id", "bbox", "purpose"], zone_context)
            if any(not isinstance(zone[field], str) or not zone[field].strip() for field in ("id", "purpose")):
                raise ContractError(f"{zone_context} id and purpose must not be blank")
            validate_fraction_bbox(zone["bbox"], f"{zone_context}.bbox")
            zone_ids.append(zone["id"])
        if len(zone_ids) != len(set(zone_ids)):
            raise ContractError(f"{context}.layout.decorative_safe_zones contains duplicate ids")


def validate_plan_docs(root: Path, project: dict[str, Any], phase: str) -> dict[str, dict[str, Any]]:
    scenario = project["scenario"]
    keys = list(PLAN_DOCS[scenario])
    if scenario == "new-deck" and phase in {"preview", "final"}:
        keys.append("style_reference_plan")
    handoff_path = root / DOCS["scene1_handoff"]
    scene1_handoff = scenario == "reconstruction" and any(
        source.get("role") == "scene1-handoff" for source in project["source_files"]
    )
    if scenario == "reconstruction" and scene1_handoff != handoff_path.is_file():
        raise ContractError(
            "scene 1 handoff requires both plans/scene1_handoff.json and "
            "project source_files with role scene1-handoff"
        )
    if scene1_handoff:
        keys.append("scene1_handoff")
    if scenario == "existing-deck" and project["new_page_count"] > 0:
        keys.extend(["outline", "slide_plan"])
    docs = {key: load_doc(root, key, project["run_id"]) for key in keys}

    brief = docs["brief"]
    require(brief, ["scenario", "audience", "objective", "success_criteria"], "brief.json")
    if brief["scenario"] != scenario:
        raise ContractError("brief.json.scenario does not match project.json")
    if phase in {"preview", "final"}:
        unconfirmed = [key for key, doc in docs.items() if doc.get("status") != CONFIRMED_STATUS]
        if unconfirmed:
            raise ContractError("preview/final requires confirmed plan files: " + ", ".join(unconfirmed))

    if "outline" in docs:
        outline_ids = slide_ids(docs["outline"], "outline.json")
        plan_ids = slide_ids(docs["slide_plan"], "slide_plan.json")
        validate_slide_plan(docs["slide_plan"])
        if outline_ids != plan_ids:
            raise ContractError("outline and slide plan roster/order differ")
        if scenario == "existing-deck" and len(plan_ids) != project["new_page_count"]:
            raise ContractError("existing-deck new slide plan count does not match project.new_page_count")

    if "style_reference_plan" in docs:
        validate_style_reference_plan(root, docs["style_reference_plan"])

    if "gorden_component" in docs:
        component = docs["gorden_component"]
        require(
            component,
            [
                "repository", "expected_commit", "component_location", "notice_path",
                "probe_report", "probe_sha256",
                "integration_mode", "pipeline", "stages", "generation_scope",
                "source_bundled", "attribution_url",
            ],
            "gorden_component.json",
        )
        if component["repository"] != "GordenSun/GordenSuperPPTSkills":
            raise ContractError("gorden component repository is invalid")
        if not re.fullmatch(r"[0-9a-f]{40}", str(component["expected_commit"])):
            raise ContractError("gorden component expected_commit must be a full Git SHA")
        if component["integration_mode"] != "bundled-attributed" or component["source_bundled"] is not True:
            raise ContractError("gorden component must remain bundled with attribution")
        if component["component_location"] != "components/gorden" or component["notice_path"] != "components/gorden/NOTICE.md":
            raise ContractError("gorden bundled component location is invalid")
        if component["attribution_url"] != "https://github.com/GordenSun/GordenSuperPPTSkills":
            raise ContractError("gorden component attribution URL is invalid")
        component_root = Path(__file__).resolve().parent.parent / "components" / "gorden"
        if not component_root.is_dir():
            raise ContractError("bundled gorden component is missing")
        if not (component_root / "NOTICE.md").is_file():
            raise ContractError("bundled gorden attribution notice is missing")
        probe_path = resolve_inside(root, component["probe_report"], "gorden probe report")
        if not probe_path.is_file() or component["probe_sha256"] != sha256(probe_path):
            raise ContractError("gorden probe report is missing or stale")
        probe = load_json(probe_path)
        gorden_probe = probe.get("gorden_super_ppt")
        if probe.get("schema") != "pptx-workshop.backend-probe.v1" or not isinstance(gorden_probe, dict):
            raise ContractError("gorden probe report schema is invalid")
        if Path(str(gorden_probe.get("root", ""))).resolve() != component_root.resolve():
            raise ContractError("gorden probe root does not match bundled component")
        if gorden_probe.get("bundled") is not True or gorden_probe.get("attribution_notice") is not True:
            raise ContractError("gorden probe must confirm bundled attribution")
        if gorden_probe.get("source_sha") != component["expected_commit"]:
            raise ContractError("gorden component checkout does not match expected_commit")
        image_ppt = gorden_probe.get("image_ppt_gen", {})
        image_to_pptx = gorden_probe.get("image_to_pptx", {})
        modules = probe.get("python_modules", {})
        if not all(modules.get(name) is True for name in ("python_pptx", "Pillow", "numpy")):
            raise ContractError("Gorden component dependencies are missing")
        expected_pipeline = "image-ppt" if scenario == "new-deck" else "image2pptx"
        expected_stages = ["GordenImagePPTGen"] if scenario == "new-deck" else ["GordenImage2PPTX"]
        stages = component["stages"]
        if component["pipeline"] != expected_pipeline or not isinstance(stages, list):
            raise ContractError(f"{scenario} gorden pipeline is invalid")
        if [item.get("component") for item in stages if isinstance(item, dict)] != expected_stages:
            raise ContractError(f"{scenario} gorden stage order is invalid")
        if scenario == "new-deck":
            if not all((image_ppt.get("available"), image_ppt.get("compose_pptx"))):
                raise ContractError("GordenImagePPTGen component is incomplete")
        else:
            required_image_to_pptx = (
                image_to_pptx.get("available"), image_to_pptx.get("compose_pptx"),
                image_to_pptx.get("layout_guard"), image_to_pptx.get("placement_qa"),
                image_to_pptx.get("visual_compare_qa"),
            )
            if not all(required_image_to_pptx):
                raise ContractError("GordenImage2PPTX component is incomplete")
        for index, stage in enumerate(stages):
            require(
                stage,
                ["component", "purpose", "output_root", "manifest_path"],
                f"gorden_component.stages[{index}]",
            )
            resolve_inside(root, stage["output_root"], "gorden stage output_root")
            resolve_inside(root, stage["manifest_path"], "gorden stage manifest_path")
        scope = component["generation_scope"]
        require(scope, ["page_count", "estimated_imagegen_calls_min", "basis", "user_notice"], "gorden generation_scope")
        page_count = len(docs["slide_plan"]["slides"]) if scenario == "new-deck" else len(docs["reconstruction_plan"]["slides"])
        representative_count = min(page_count, 3)
        minimum_calls = (
            page_count + 2 * representative_count
            if scenario == "new-deck"
            else page_count * 3
        )
        if scope["page_count"] != page_count:
            raise ContractError("gorden generation_scope.page_count does not match the plan")
        if scope["estimated_imagegen_calls_min"] < minimum_calls:
            raise ContractError("gorden generation_scope underestimates minimum imagegen calls")
        if not str(scope["basis"]).strip() or not str(scope["user_notice"]).strip():
            raise ContractError("gorden generation scope must explain its estimate and user notice")

    if scenario == "new-deck":
        planned_slide_ids = slide_ids(docs["slide_plan"], "slide_plan.json")
        assets = docs["asset_plan"].get("items")
        if not isinstance(assets, list):
            raise ContractError("asset_plan.items must be an array")
        final_images: dict[str, dict[str, Any]] = {}
        for index, item in enumerate(assets):
            if not isinstance(item, dict) or item.get("visual_role") != "full-slide-image":
                continue
            context = f"asset_plan.items[{index}]"
            require(
                item,
                [
                    "id", "pages", "purpose", "asset_type", "source",
                    "contains_semantic_text", "ordinary_text_in_asset", "visual_role",
                    "source_kind", "selection_reason", "editability_impact", "bbox",
                    "pixel_width", "pixel_height",
                ],
                context,
            )
            pages = item["pages"]
            if not isinstance(pages, list) or len(pages) != 1 or pages[0] not in planned_slide_ids:
                raise ContractError(f"{context}.pages must name exactly one planned slide")
            page_id = pages[0]
            if page_id in final_images:
                raise ContractError(f"new-deck has multiple full-slide images for {page_id}")
            if (
                item["asset_type"] != "png"
                or not str(item["source"]).lower().endswith(".png")
                or item["source_kind"] != "generated"
                or item["contains_semantic_text"] is not True
                or item["ordinary_text_in_asset"] is not True
                or item["bbox"] != [0, 0, 1, 1]
                or not str(item["selection_reason"]).strip()
                or not str(item["editability_impact"]).strip()
            ):
                raise ContractError(f"{context} must describe one generated, text-bearing, full-slide PNG")
            if (
                not isinstance(item["pixel_width"], int)
                or not isinstance(item["pixel_height"], int)
                or item["pixel_width"] < 1
                or item["pixel_height"] < 1
                or max(item["pixel_width"], item["pixel_height"]) < 1920
            ):
                raise ContractError(f"{context} must declare a full-slide PNG with long edge at least 1920 px")
            if phase == "final":
                image_path = resolve_inside(root, item["source"], "new-deck final PNG")
                if not image_path.is_file():
                    raise ContractError(f"new-deck final PNG is missing: {image_path}")
                if png_dimensions(image_path) != (item["pixel_width"], item["pixel_height"]):
                    raise ContractError(f"new-deck final PNG dimensions do not match asset plan: {item['source']}")
            final_images[page_id] = item
        if list(final_images) != planned_slide_ids:
            raise ContractError("new-deck asset plan must contain exactly one ordered full-slide PNG for every slide")

    if scenario == "template-fill":
        analysis = docs["template_analysis"]
        if analysis.get("template_kind") not in {"placeholder-based", "framework-only", "mixed"}:
            raise ContractError("template_analysis.template_kind is invalid")
        if not isinstance(analysis.get("slide_families"), list) or not analysis["slide_families"]:
            raise ContractError("template_analysis.slide_families must not be empty")
        if not slide_ids(docs["fill_plan"], "fill_plan.json"):
            raise ContractError("fill plan must contain slides")

    if scenario == "reconstruction":
        reconstruction = docs["reconstruction_plan"]
        if reconstruction.get("source_locked") is not True:
            raise ContractError("reconstruction_plan.source_locked must be true")
        if reconstruction.get("visual_asset_policy") != "fidelity-first-controlled-png":
            raise ContractError("reconstruction_plan.visual_asset_policy must be fidelity-first-controlled-png")
        ids = slide_ids(reconstruction, "reconstruction_plan.json")
        pages = [slide.get("source_page") for slide in reconstruction["slides"]]
        if pages != list(range(1, len(pages) + 1)):
            raise ContractError("reconstruction source pages must stay sequential and start at 1")
        for slide in reconstruction["slides"]:
            for item in slide.get("objects", []):
                if item.get("semantic_role") == "text" and item.get("target_kind") != "native_text":
                    raise ContractError(f"{slide['id']} semantic text must target native_text")
        assets = docs["asset_plan"].get("items")
        if not isinstance(assets, list):
            raise ContractError("asset_plan.items must be an array")
        page_map = {slide["id"]: slide for slide in reconstruction["slides"]}
        asset_ids: set[str] = set()
        for index, item in enumerate(assets):
            context = f"asset_plan.items[{index}]"
            if not isinstance(item, dict):
                raise ContractError(f"{context} must be an object")
            require(
                item,
                [
                    "id", "pages", "purpose", "asset_type", "source",
                    "contains_semantic_text", "ordinary_text_in_asset",
                    "semantic_text_policy", "text_removal_status", "visual_role",
                    "source_kind", "selection_reason", "editability_impact", "bbox",
                ],
                context,
            )
            if item["id"] in asset_ids:
                raise ContractError(f"duplicate reconstruction asset id: {item['id']}")
            asset_ids.add(item["id"])
            if item["asset_type"] not in {"svg", "png"}:
                raise ContractError("reconstruction asset_type must be svg or png")
            if item["visual_role"] not in {
                "background", "frame", "photo", "texture", "illustration", "screenshot",
                "chart-or-diagram", "icon", "artistic-text", "other-complex-visual",
            }:
                raise ContractError(f"{context}.visual_role is invalid")
            if item["source_kind"] not in {"original", "extracted", "generated", "redrawn"}:
                raise ContractError(f"{context}.source_kind is invalid")
            if not str(item["selection_reason"]).strip() or not str(item["editability_impact"]).strip():
                raise ContractError(f"{context} must explain fidelity choice and editability impact")
            item_pages = item["pages"]
            if not isinstance(item_pages, list) or not item_pages or any(page not in page_map for page in item_pages):
                raise ContractError(f"{context}.pages must reference reconstruction slide ids")
            bbox = item["bbox"]
            if (
                not isinstance(bbox, list)
                or len(bbox) != 4
                or any(not isinstance(value, (int, float)) for value in bbox)
                or bbox[0] < 0
                or bbox[1] < 0
                or bbox[2] <= 0
                or bbox[3] <= 0
                or bbox[0] + bbox[2] > 1.000001
                or bbox[1] + bbox[3] > 1.000001
            ):
                raise ContractError(f"{context}.bbox must be a positive fraction bbox inside the slide")
            if item["ordinary_text_in_asset"] is not False:
                raise ContractError("reconstruction visual assets may not contain ordinary text")
            if item["contains_semantic_text"] is True:
                if (
                    item["visual_role"] != "artistic-text"
                    or item["semantic_text_policy"] != "artistic-user-confirmation"
                    or item["text_removal_status"] != "artistic-only-user-confirmation"
                ):
                    raise ContractError("semantic text in a visual asset is allowed only as separately confirmed artistic text")
            elif item["semantic_text_policy"] != "none":
                raise ContractError("non-text reconstruction assets must use semantic_text_policy=none")
            source = str(item["source"])
            if item["asset_type"] == "svg":
                if not source.lower().endswith(".svg"):
                    raise ContractError("reconstruction SVG asset source must end in .svg")
            else:
                if not source.lower().endswith(".png"):
                    raise ContractError("reconstruction raster assets must be lossless .png files")
                if item["text_removal_status"] not in {"verified-clean", "artistic-only-user-confirmation"}:
                    raise ContractError("reconstruction PNG must have verified text removal or approved artistic text")
                if not isinstance(item.get("pixel_width"), int) or not isinstance(item.get("pixel_height"), int):
                    raise ContractError("reconstruction PNG must declare pixel_width and pixel_height")
                if item["pixel_width"] < 1 or item["pixel_height"] < 1:
                    raise ContractError("reconstruction PNG dimensions must be positive")
                if phase in {"preview", "final"}:
                    png_path = resolve_inside(root, source, "reconstruction PNG source")
                    if not png_path.is_file():
                        raise ContractError(f"reconstruction PNG is missing: {png_path}")
                    actual_width, actual_height = png_dimensions(png_path)
                    if (actual_width, actual_height) != (item["pixel_width"], item["pixel_height"]):
                        raise ContractError(f"reconstruction PNG dimensions do not match asset plan: {source}")
                    for page_id in item_pages:
                        page = page_map[page_id]
                        scale = min(1.0, 1920 / max(page["ref_width"], page["ref_height"]))
                        minimum_width = math.ceil(page["ref_width"] * scale * bbox[2])
                        minimum_height = math.ceil(page["ref_height"] * scale * bbox[3])
                        if actual_width < minimum_width or actual_height < minimum_height:
                            raise ContractError(
                                f"reconstruction PNG resolution is too low for {page_id}: "
                                f"{actual_width}x{actual_height} < {minimum_width}x{minimum_height}"
                            )
        if ids != [f"P{i:02d}" for i in range(1, len(ids) + 1)]:
            raise ContractError("reconstruction slide ids must preserve source order")

        if scene1_handoff:
            handoff = docs["scene1_handoff"]
            require(
                handoff,
                [
                    "upstream_run_id", "upstream_project", "upstream_artifact_manifest",
                    "upstream_asset_plan", "upstream_final_pptx",
                    "upstream_image_deck_approval", "pages",
                ],
                "scene1_handoff.json",
            )
            upstream_project_ref = handoff["upstream_project"]
            upstream_manifest_ref = handoff["upstream_artifact_manifest"]
            upstream_asset_plan_ref = handoff["upstream_asset_plan"]
            upstream_final = handoff["upstream_final_pptx"]
            upstream_approval = handoff["upstream_image_deck_approval"]
            require(upstream_project_ref, ["path", "sha256"], "scene1_handoff.upstream_project")
            require(
                upstream_manifest_ref,
                ["path", "sha256"],
                "scene1_handoff.upstream_artifact_manifest",
            )
            require(
                upstream_asset_plan_ref,
                ["path", "sha256"],
                "scene1_handoff.upstream_asset_plan",
            )
            require(upstream_final, ["path", "sha256"], "scene1_handoff.upstream_final_pptx")
            require(upstream_approval, ["path", "sha256"], "scene1_handoff.upstream_image_deck_approval")
            upstream_project_path = Path(upstream_project_ref["path"]).expanduser().resolve()
            upstream_manifest_path = Path(upstream_manifest_ref["path"]).expanduser().resolve()
            upstream_asset_plan_path = Path(upstream_asset_plan_ref["path"]).expanduser().resolve()
            final_path = Path(upstream_final["path"]).expanduser().resolve()
            approval_path = Path(upstream_approval["path"]).expanduser().resolve()
            referenced_files = (
                (upstream_project_path, upstream_project_ref["sha256"], "project"),
                (upstream_manifest_path, upstream_manifest_ref["sha256"], "artifact manifest"),
                (upstream_asset_plan_path, upstream_asset_plan_ref["sha256"], "asset plan"),
            )
            for referenced_path, expected_hash, label in referenced_files:
                if not referenced_path.is_file() or sha256(referenced_path) != expected_hash:
                    raise ContractError(f"scene 1 handoff upstream {label} is missing or stale")
            if not final_path.is_file() or sha256(final_path) != upstream_final["sha256"]:
                raise ContractError("scene 1 handoff final PPTX is missing or stale")
            if not approval_path.is_file() or sha256(approval_path) != upstream_approval["sha256"]:
                raise ContractError("scene 1 handoff image-deck-final approval is missing or stale")
            upstream_root = upstream_project_path.parent
            if upstream_root == root or upstream_project_path != upstream_root / "project.json":
                raise ContractError("scene 1 handoff must reference a separate upstream run directory")
            upstream_project = load_json(upstream_project_path)
            if (
                upstream_project.get("schema_version") != "pptx-workshop.project.v1"
                or upstream_project.get("scenario") != "new-deck"
                or upstream_project.get("run_id") != handoff["upstream_run_id"]
            ):
                raise ContractError("scene 1 handoff upstream project/run identity is invalid")
            if upstream_manifest_path != upstream_root / DOCS["artifacts"]:
                raise ContractError("scene 1 handoff artifact manifest path is invalid")
            upstream_manifest = load_json(upstream_manifest_path)
            if (
                upstream_manifest.get("schema_version") != SCHEMA_VERSIONS["artifacts"]
                or upstream_manifest.get("run_id") != handoff["upstream_run_id"]
            ):
                raise ContractError("scene 1 handoff artifact manifest identity is invalid")
            upstream_finals = artifact_by_purpose(upstream_manifest.get("artifacts", []), "final-output")
            if len(upstream_finals) != 1:
                raise ContractError("scene 1 handoff upstream run must have exactly one final-output PPTX")
            upstream_final_artifact = upstream_finals[0]
            expected_final_path = resolve_inside(
                upstream_root,
                upstream_final_artifact.get("path", ""),
                "scene 1 upstream final artifact",
            )
            if (
                expected_final_path != final_path
                or upstream_final_artifact.get("sha256") != upstream_final["sha256"]
            ):
                raise ContractError("scene 1 handoff final PPTX does not match the upstream artifact manifest")
            if upstream_asset_plan_path != upstream_root / DOCS["asset_plan"]:
                raise ContractError("scene 1 handoff asset plan path is invalid")
            upstream_asset_plan = load_json(upstream_asset_plan_path)
            if (
                upstream_asset_plan.get("schema_version") != SCHEMA_VERSIONS["asset_plan"]
                or upstream_asset_plan.get("run_id") != handoff["upstream_run_id"]
                or upstream_asset_plan.get("status") != CONFIRMED_STATUS
            ):
                raise ContractError("scene 1 handoff upstream asset plan identity is invalid")
            approval = load_json(approval_path)
            if (
                approval.get("gate") != "image-deck-final"
                or approval.get("status") != "confirmed"
                or approval.get("confirmed_by") != "user"
            ):
                raise ContractError("scene 1 handoff must reference a confirmed image-deck-final approval")
            if approval_path.parent.parent != upstream_root:
                raise ContractError("scene 1 handoff approval must belong to the upstream run")
            approved_subject = (upstream_root / approval.get("subject_path", "")).resolve()
            if approved_subject != final_path or approval.get("subject_sha256") != upstream_final["sha256"]:
                raise ContractError("scene 1 handoff approval does not bind the upstream final PPTX")
            pages = handoff["pages"]
            if not isinstance(pages, list) or [page.get("id") for page in pages if isinstance(page, dict)] != ids:
                raise ContractError("scene 1 handoff pages must match reconstruction slide order")
            upstream_assets = [
                item
                for item in upstream_asset_plan.get("items", [])
                if isinstance(item, dict) and item.get("visual_role") == "full-slide-image"
            ]
            upstream_asset_ids = [
                item.get("pages", [None])[0]
                for item in upstream_assets
                if isinstance(item.get("pages"), list) and len(item["pages"]) == 1
            ]
            if upstream_asset_ids != ids:
                raise ContractError("scene 1 handoff must cover every upstream final page in order")
            embedded_hashes = inspect_image_deck_pptx(final_path)
            if len(embedded_hashes) != len(pages):
                raise ContractError("scene 1 handoff page count differs from the approved image deck")
            handoff_sources: list[tuple[Path, str]] = []
            for index, (page, planned_asset) in enumerate(zip(pages, upstream_assets)):
                require(page, ["id", "path", "sha256"], f"scene1_handoff.pages[{index}]")
                page_path = Path(page["path"]).expanduser().resolve()
                if not page_path.is_file() or sha256(page_path) != page["sha256"]:
                    raise ContractError(f"scene 1 handoff page is missing or stale: {page_path}")
                png_dimensions(page_path)
                planned_path = resolve_inside(
                    upstream_root,
                    planned_asset.get("source", ""),
                    "scene 1 upstream page asset",
                )
                if page_path != planned_path or embedded_hashes[index] != page["sha256"]:
                    raise ContractError(
                        "scene 1 handoff pages must be the ordered PNGs embedded in the approved image deck"
                    )
                handoff_sources.append((page_path, page["sha256"]))
            project_sources = [
                (Path(source["path"]).expanduser().resolve(), source["sha256"])
                for source in project["source_files"]
                if source.get("role") == "scene1-handoff"
            ]
            if project_sources != handoff_sources:
                raise ContractError("project scene1-handoff sources must exactly match the handoff pages")

    if scenario == "existing-deck":
        change = docs["change_plan"]
        if not isinstance(change.get("changes"), list) or not isinstance(change.get("new_pages"), list):
            raise ContractError("change_plan changes/new_pages must be arrays")
        if len(change["new_pages"]) != project["new_page_count"]:
            raise ContractError("change_plan.new_pages count does not match project.json")
        for index, item in enumerate(change["changes"]):
            require(
                item,
                ["source_slide", "selector", "before", "after", "preserve"],
                f"change_plan.changes[{index}]",
            )
            required_format = {"font", "font_size", "color", "position", "size", "paragraph", "fill", "border"}
            if not required_format.issubset(set(item["preserve"])):
                raise ContractError("existing-deck changes must preserve the complete format contract")
    return docs


def approval_records(root: Path) -> list[dict[str, Any]]:
    records: list[dict[str, Any]] = []
    approval_dir = root / "approvals"
    if not approval_dir.exists():
        return records
    for path in sorted(approval_dir.glob("*.json")):
        record = load_json(path)
        require(
            record,
            ["schema_version", "gate", "status", "confirmed_by", "subject_path", "subject_sha256", "confirmed_at", "decision"],
            str(path.relative_to(root)),
        )
        if record["schema_version"] != "pptx-workshop.approval.v1":
            raise ContractError(f"invalid approval schema: {path}")
        if record["status"] != "confirmed" or record["confirmed_by"] != "user":
            raise ContractError(f"approval must be explicitly confirmed by user: {path}")
        subject = resolve_inside(root, record["subject_path"], "approval subject")
        if not subject.is_file() or record["subject_sha256"] != sha256(subject):
            raise ContractError(f"approval subject is missing or stale: {path}")
        records.append(record)
    return records


def require_approval(
    records: list[dict[str, Any]], gate: str, root: Path, subject: Path | None = None
) -> dict[str, Any]:
    candidates = [record for record in records if record["gate"] == gate]
    if subject is not None:
        relative = subject.resolve().relative_to(root).as_posix()
        candidates = [record for record in candidates if Path(record["subject_path"]).as_posix() == relative]
    if not candidates:
        label = f" for {subject.relative_to(root)}" if subject else ""
        raise ContractError(f"missing user approval: {gate}{label}")
    return candidates[0]


def pptx_slide_count(path: Path) -> int:
    try:
        with zipfile.ZipFile(path) as archive:
            return sum(
                1
                for name in archive.namelist()
                if re.fullmatch(r"ppt/slides/slide\d+\.xml", name)
            )
    except (OSError, zipfile.BadZipFile) as exc:
        raise ContractError(f"invalid PPTX package: {path}: {exc}") from exc


def inspect_image_deck_pptx(path: Path) -> list[str]:
    try:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE

        presentation = Presentation(path)
    except Exception as exc:
        raise ContractError(f"cannot read scene 1 image PPTX: {path}: {exc}") from exc
    image_hashes: list[str] = []
    for index, slide in enumerate(presentation.slides, 1):
        shapes = list(slide.shapes)
        if len(shapes) != 1 or shapes[0].shape_type != MSO_SHAPE_TYPE.PICTURE:
            raise ContractError(f"scene 1 PPTX slide {index} must contain exactly one picture object")
        picture = shapes[0]
        if (
            picture.left != 0
            or picture.top != 0
            or picture.width != presentation.slide_width
            or picture.height != presentation.slide_height
        ):
            raise ContractError(f"scene 1 PPTX slide {index} picture must cover the full canvas")
        if picture.image.ext.lower() != "png":
            raise ContractError(f"scene 1 PPTX slide {index} embedded picture must be PNG")
        image_hashes.append(hashlib.sha256(picture.image.blob).hexdigest())
    return image_hashes


def validate_image_deck_pptx(path: Path, root: Path, assets: list[dict[str, Any]]) -> None:
    image_hashes = inspect_image_deck_pptx(path)
    if len(image_hashes) != len(assets):
        raise ContractError("scene 1 image PPTX page count differs from its planned PNG set")
    for index, (image_hash, asset) in enumerate(zip(image_hashes, assets), 1):
        source = resolve_inside(root, asset["source"], "scene 1 planned PNG")
        if image_hash != sha256(source):
            raise ContractError(f"scene 1 PPTX slide {index} does not embed its confirmed PNG")


def require_artifact_parent(root: Path, artifact: dict[str, Any], required_path: Path) -> None:
    relative = required_path.resolve().relative_to(root).as_posix()
    current_hash = sha256(required_path)
    if not any(
        Path(parent.get("path", "")).as_posix() == relative
        and parent.get("sha256") == current_hash
        for parent in artifact.get("parents", [])
    ):
        raise ContractError(f"{artifact.get('purpose', 'artifact')} does not bind required parent: {relative}")


def validate_gorden_generation_manifest(
    root: Path,
    project: dict[str, Any],
    docs: dict[str, dict[str, Any]],
    artifacts: list[dict[str, Any]],
    approvals: list[dict[str, Any]],
    phase: str,
) -> None:
    component = docs.get("gorden_component")
    if component is None or phase == "plan":
        return
    stages = component.get("stages", [])
    if len(stages) != 1:
        raise ContractError("Gorden generation manifest requires exactly one active stage")
    manifest_path = resolve_inside(root, stages[0]["manifest_path"], "Gorden generation manifest")
    manifest = load_json(manifest_path)
    require(
        manifest,
        ["schema_version", "run_id", "component", "calls", "qa_results"],
        "Gorden generation manifest",
    )
    expected_component = (
        "GordenImagePPTGen" if project["scenario"] == "new-deck" else "GordenImage2PPTX"
    )
    if (
        manifest["schema_version"] != GORDEN_GENERATION_MANIFEST_VERSION
        or manifest["run_id"] != project["run_id"]
        or manifest["component"] != expected_component
    ):
        raise ContractError("Gorden generation manifest identity is invalid")
    calls = manifest["calls"]
    if not isinstance(calls, list) or not calls:
        raise ContractError("Gorden generation manifest calls must not be empty")
    if not isinstance(manifest["qa_results"], list):
        raise ContractError("Gorden generation manifest qa_results must be an array")

    page_ids = (
        slide_ids(docs["slide_plan"], "slide_plan.json")
        if project["scenario"] == "new-deck"
        else slide_ids(docs["reconstruction_plan"], "reconstruction_plan.json")
    )
    page_id_set = set(page_ids)
    scope_approval = require_approval(
        approvals,
        "gorden-generation-scope",
        root,
        root / DOCS["gorden_component"],
    )
    scope_time = parse_timestamp(scope_approval["confirmed_at"], "gorden-generation-scope.confirmed_at")
    slide_map: dict[str, dict[str, Any]] = {}
    representative_ids: list[str] = []
    style_reference_ids: list[str] = []
    style_reference_choice_time: datetime | None = None
    style_choice: dict[str, Any] | None = None
    style_choice_time: datetime | None = None
    if project["scenario"] == "new-deck":
        slide_map = {slide["id"]: slide for slide in docs["slide_plan"]["slides"]}
        style_reference_plan = docs["style_reference_plan"]
        style_reference_ids = style_reference_plan["selected_ids"]
        style_reference_choice = require_approval(
            approvals,
            "style-reference-choice",
            root,
            root / DOCS["style_reference_plan"],
        )
        style_reference_decision = style_reference_choice.get("decision", {})
        if (
            style_reference_decision.get("selected_ids") != style_reference_ids
            or style_reference_decision.get("usage_mode") != style_reference_plan["usage_mode"]
        ):
            raise ContractError("style-reference-choice differs from style_reference_plan.json")
        style_reference_choice_time = parse_timestamp(
            style_reference_choice["confirmed_at"],
            "style-reference-choice.confirmed_at",
        )
        style_choice = require_approval(approvals, "style-choice", root)
        style_choice_time = parse_timestamp(style_choice["confirmed_at"], "style-choice.confirmed_at")
        representatives = style_choice.get("decision", {}).get("representatives")
        if not isinstance(representatives, list):
            raise ContractError("style-choice must list reusable representative pages")
        representative_ids = [
            item.get("slide_id") for item in representatives if isinstance(item, dict)
        ]
        if len(representative_ids) != len(representatives):
            raise ContractError("style-choice representative page mapping is invalid")

    call_ids: set[str] = set()
    successful: dict[tuple[str, str], list[str]] = {}
    attempted: dict[str, set[str]] = {}
    for index, call in enumerate(calls):
        if not isinstance(call, dict):
            raise ContractError(f"Gorden generation manifest calls[{index}] must be an object")
        require(
            call,
            [
                "id",
                "purpose",
                "workflow_phase",
                "page_id",
                "provider",
                "model",
                "prompt",
                "generated_at",
                "status",
            ],
            f"Gorden generation manifest calls[{index}]",
        )
        if call["id"] in call_ids:
            raise ContractError("Gorden generation manifest contains duplicate call ids")
        call_ids.add(call["id"])
        if any(
            not str(call[field]).strip()
            for field in (
                "purpose",
                "workflow_phase",
                "page_id",
                "provider",
                "model",
                "prompt",
                "generated_at",
            )
        ):
            raise ContractError(f"Gorden generation manifest calls[{index}] has blank provenance fields")
        if call["page_id"] not in page_id_set:
            raise ContractError(f"Gorden generation manifest call targets an unplanned page: {call['page_id']}")
        if call["status"] not in {"succeeded", "failed"}:
            raise ContractError(f"Gorden generation manifest calls[{index}].status is invalid")
        generated_at = parse_timestamp(
            call["generated_at"], f"Gorden generation manifest calls[{index}].generated_at"
        )
        purpose = call["purpose"]
        attempted.setdefault(purpose, set()).add(call["page_id"])

        if project["scenario"] == "new-deck":
            expected_phases = {
                "style-option-a": "style-selection",
                "style-option-b": "style-selection",
                "style-option-c": "style-selection",
                "final-page": "formal-generation",
            }
            if purpose not in expected_phases or call["workflow_phase"] != expected_phases[purpose]:
                raise ContractError("scene 1 calls must separate style-selection from formal-generation")
            require(
                call,
                ["slide_contract_sha256", "style_reference_ids"],
                f"Gorden generation manifest calls[{index}]",
            )
            if call["style_reference_ids"] != style_reference_ids:
                raise ContractError("scene 1 generation call is not bound to the confirmed style references")
            slide = slide_map[call["page_id"]]
            if call["slide_contract_sha256"] != json_sha256(slide):
                raise ContractError("scene 1 generation call is not bound to the current slide contract")
            prompt = call["prompt"]
            for block in slide["content_blocks"]:
                compact_bbox = json.dumps(block["bbox"], ensure_ascii=False, separators=(",", ":"))
                spaced_bbox = json.dumps(block["bbox"], ensure_ascii=False)
                if block["text"] not in prompt or (
                    compact_bbox not in prompt and spaced_bbox not in prompt
                ):
                    raise ContractError(
                        "scene 1 prompt must contain every locked text block and its bbox"
                    )
            if generated_at < scope_time:
                raise ContractError("scene 1 image generation occurred before scope confirmation")
            if style_reference_choice_time is None or generated_at < style_reference_choice_time:
                raise ContractError("scene 1 image generation occurred before style-reference-choice")
            if purpose.startswith("style-option-"):
                if style_choice_time is None or generated_at > style_choice_time:
                    raise ContractError("style-option images must be generated before style-choice")
            elif style_choice_time is None or generated_at < style_choice_time:
                raise ContractError("formal-generation may start only after style-choice")
        else:
            if purpose not in {"background", "frame", "icons"} or call["workflow_phase"] != "reconstruction":
                raise ContractError("reconstruction calls must use the reconstruction workflow phase")

        if call["status"] == "failed":
            continue
        require(
            call,
            ["output_path", "output_sha256"],
            f"Gorden generation manifest calls[{index}]",
        )
        output_path = resolve_inside(root, call["output_path"], "Gorden generated output")
        if not output_path.is_file() or call["output_sha256"] != sha256(output_path):
            raise ContractError("Gorden generation manifest output is missing or stale")
        png_dimensions(output_path)
        successful.setdefault((purpose, call["page_id"]), []).append(call["output_sha256"])

    if project["scenario"] == "new-deck":
        representative_count = min(len(page_ids), 3)
        if len(representative_ids) != representative_count or len(set(representative_ids)) != representative_count:
            raise ContractError("style-choice must identify every distinct reusable representative page")
        representative_set = set(representative_ids)
        if len(calls) < 3 * representative_count:
            raise ContractError("scene 1 generation manifest undercounts style-option imagegen calls")
        for letter in ("a", "b", "c"):
            purpose = f"style-option-{letter}"
            attempted_pages = attempted.get(purpose, set())
            successful_pages = {
                page_id for call_purpose, page_id in successful if call_purpose == purpose
            }
            if attempted_pages != representative_set or successful_pages != representative_set:
                raise ContractError(
                    f"{purpose} must generate the same confirmed representative page set"
                )
            options = artifact_by_purpose(artifacts, purpose)
            if len(options) != 1:
                raise ContractError(f"generation manifest cannot resolve {purpose}")
            hashes = inspect_image_deck_pptx(artifact_path(root, options[0]))
            if len(hashes) != representative_count or any(
                digest not in successful.get((purpose, page_id), [])
                for page_id, digest in zip(representative_ids, hashes)
            ):
                raise ContractError(
                    f"{purpose} page order or image provenance differs from the representative mapping"
                )
        remaining_set = page_id_set - representative_set
        attempted_final_pages = attempted.get("final-page", set())
        if not attempted_final_pages.issubset(remaining_set):
            raise ContractError("selected representative pages may not be regenerated in formal-generation")
        if phase == "final":
            minimum_calls = len(page_ids) + 2 * representative_count
            if len(calls) < minimum_calls:
                raise ContractError("scene 1 generation manifest undercounts final imagegen calls")
            successful_final_pages = {
                page_id for call_purpose, page_id in successful if call_purpose == "final-page"
            }
            if successful_final_pages != remaining_set:
                raise ContractError("formal-generation must successfully generate every remaining page exactly by role")
            finals = artifact_by_purpose(artifacts, "final-output")
            if len(finals) != 1:
                raise ContractError("scene 1 final deck is missing from the generation manifest check")
            selected_paths = {
                Path(artifact["path"]).as_posix(): artifact["purpose"]
                for artifact in artifacts
                if artifact.get("purpose") in {
                    "style-option-a",
                    "style-option-b",
                    "style-option-c",
                }
            }
            assert style_choice is not None
            selected_purpose = selected_paths.get(Path(style_choice["subject_path"]).as_posix())
            if selected_purpose is None:
                raise ContractError("style-choice does not identify a generated style option")
            final_hashes = inspect_image_deck_pptx(artifact_path(root, finals[0]))
            if len(final_hashes) != len(page_ids):
                raise ContractError("scene 1 final deck page count differs from the plan")
            for page_id, digest in zip(page_ids, final_hashes):
                expected_purpose = selected_purpose if page_id in representative_set else "final-page"
                if digest not in successful.get((expected_purpose, page_id), []):
                    raise ContractError(
                        "scene 1 final deck must reuse selected representatives and generate only remaining pages"
                    )
    else:
        required_purposes = {"background", "frame", "icons"}
        successful_purposes = {purpose for purpose, _page_id in successful}
        if not required_purposes.issubset(successful_purposes):
            raise ContractError("reconstruction manifest must record background, frame and icons generation")
        if phase == "final":
            if len(calls) < len(page_ids) * 3:
                raise ContractError("reconstruction generation manifest undercounts final imagegen calls")
            for page_id in page_ids:
                for purpose in required_purposes:
                    if not successful.get((purpose, page_id)):
                        raise ContractError(
                            f"reconstruction generation manifest is missing {purpose} for {page_id}"
                        )
        qa_results = manifest["qa_results"]
        qa_tools: set[str] = set()
        for index, result in enumerate(qa_results):
            if not isinstance(result, dict):
                raise ContractError(f"Gorden generation manifest qa_results[{index}] must be an object")
            require(
                result,
                ["tool", "status", "report_path", "report_sha256"],
                f"Gorden generation manifest qa_results[{index}]",
            )
            if result["tool"] in qa_tools:
                raise ContractError("Gorden generation manifest contains duplicate QA tools")
            qa_tools.add(result["tool"])
            report_path = resolve_inside(root, result["report_path"], "Gorden QA report")
            if (
                result["status"] != "pass"
                or not report_path.is_file()
                or result["report_sha256"] != sha256(report_path)
            ):
                raise ContractError(f"Gorden QA result is missing, stale or failed: {result['tool']}")
        if qa_tools != {"layout_guard", "placement_qa", "visual_compare_qa"}:
            raise ContractError("reconstruction requires all three Gorden QA results")


def validate_artifacts(root: Path, run_id: str) -> tuple[dict[str, Any] | None, list[dict[str, Any]]]:
    actual = {
        path.relative_to(root).as_posix()
        for path in root.rglob("*.pptx")
        if path.is_file()
    }
    manifest_path = root / DOCS["artifacts"]
    if not actual and not manifest_path.exists():
        return None, []
    manifest = load_doc(root, "artifacts", run_id)
    artifacts = manifest.get("artifacts")
    if not isinstance(artifacts, list) or not artifacts:
        raise ContractError("pptx_artifacts.json.artifacts must not be empty")
    registered: set[str] = set()
    ids: set[str] = set()
    for index, artifact in enumerate(artifacts):
        context = f"pptx_artifacts.artifacts[{index}]"
        require(
            artifact,
            ["id", "purpose", "path", "sha256", "qa_path", "expected_pages", "reusable", "parents"],
            context,
        )
        if artifact["id"] in ids:
            raise ContractError(f"duplicate artifact id: {artifact['id']}")
        ids.add(artifact["id"])
        pptx = resolve_inside(root, artifact["path"], context)
        if not pptx.is_file() or pptx.suffix.lower() != ".pptx":
            raise ContractError(f"artifact PPTX missing: {pptx}")
        relative = pptx.relative_to(root).as_posix()
        registered.add(relative)
        if artifact["sha256"] != sha256(pptx):
            raise ContractError(f"artifact hash mismatch: {relative}")
        expected_pages = artifact["expected_pages"]
        if not isinstance(expected_pages, int) or expected_pages < 1:
            raise ContractError(f"{context}.expected_pages must be positive")
        if pptx_slide_count(pptx) != expected_pages:
            raise ContractError(f"PPTX slide count does not match expected_pages: {relative}")
        parents = artifact["parents"]
        if not isinstance(parents, list) or not parents:
            raise ContractError(f"{context}.parents must be a non-empty array")
        for parent_index, parent in enumerate(parents):
            parent_context = f"{context}.parents[{parent_index}]"
            require(parent, ["path", "sha256"], parent_context)
            parent_path = resolve_inside(root, parent["path"], parent_context)
            if not parent_path.is_file() or parent["sha256"] != sha256(parent_path):
                raise ContractError(f"artifact parent is missing or stale: {parent['path']}")

        qa_path = resolve_inside(root, artifact["qa_path"], f"{context}.qa_path")
        qa = load_json(qa_path)
        require(
            qa,
            ["schema_version", "artifact_id", "pptx_path", "pptx_sha256", "render_dir", "rendered_pages", "inspected_pages", "inspected_by", "status", "checks", "issues"],
            str(qa_path.relative_to(root)),
        )
        if qa["schema_version"] != "pptx-workshop.visual-qa.v1":
            raise ContractError(f"invalid visual QA schema: {qa_path}")
        if qa["artifact_id"] != artifact["id"] or qa["pptx_sha256"] != artifact["sha256"]:
            raise ContractError(f"visual QA does not bind current artifact: {qa_path}")
        if Path(qa["pptx_path"]).as_posix() != Path(artifact["path"]).as_posix():
            raise ContractError(f"visual QA pptx_path mismatch: {qa_path}")
        if qa["status"] != "pass" or qa["inspected_by"] not in {"agent", "user"}:
            raise ContractError(f"visual QA must be passed after actual inspection: {qa_path}")
        if not isinstance(qa["issues"], list) or qa["issues"]:
            raise ContractError(f"passed visual QA must have an empty issues list: {qa_path}")
        render_dir = resolve_inside(root, qa["render_dir"], "visual QA render_dir")
        if not render_dir.is_dir():
            raise ContractError(f"visual QA render_dir missing: {render_dir}")
        rendered = qa["rendered_pages"]
        inspected = qa["inspected_pages"]
        if len(rendered) != expected_pages or inspected != list(range(1, expected_pages + 1)):
            raise ContractError(f"visual QA must render and inspect every page: {qa_path}")
        for page in rendered:
            image = resolve_inside(root, page, "visual QA rendered page")
            if not image.is_file() or image.suffix.lower() != ".png":
                raise ContractError(f"rendered PNG page missing: {image}")
            png_dimensions(image)
        if len(set(rendered)) != len(rendered):
            raise ContractError(f"visual QA rendered page paths must be unique: {qa_path}")
        diagnostics = qa.get("diagnostics", [])
        if not isinstance(diagnostics, list):
            raise ContractError(f"visual QA diagnostics must be an array: {qa_path}")
        for diagnostic in diagnostics:
            diagnostic_path = resolve_inside(root, diagnostic, "visual QA diagnostic")
            if not diagnostic_path.is_file() or diagnostic_path.suffix.lower() != ".png":
                raise ContractError(f"visual QA diagnostic PNG missing: {diagnostic_path}")
            png_dimensions(diagnostic_path)
    if registered != actual:
        missing = sorted(actual - registered)
        stale = sorted(registered - actual)
        raise ContractError(f"PPTX registry mismatch; unregistered={missing}, missing={stale}")
    return manifest, artifacts


def artifact_by_purpose(artifacts: list[dict[str, Any]], purpose: str) -> list[dict[str, Any]]:
    return [artifact for artifact in artifacts if artifact.get("purpose") == purpose]


def artifact_path(root: Path, artifact: dict[str, Any]) -> Path:
    return resolve_inside(root, artifact["path"], "artifact path")


def require_preview_contract(
    root: Path,
    project: dict[str, Any],
    docs: dict[str, dict[str, Any]],
    artifacts: list[dict[str, Any]],
    approvals: list[dict[str, Any]],
) -> None:
    scenario = project["scenario"]
    if not artifacts:
        raise ContractError("preview requires registered PPTX artifacts")

    if "evidence_plan" in docs and any(item.get("included") for item in docs["evidence_plan"].get("items", [])):
        require_approval(approvals, "evidence-inclusion", root, root / DOCS["evidence_plan"])
    if "asset_plan" in docs:
        require_approval(approvals, "asset-plan", root, root / DOCS["asset_plan"])
    if "gorden_component" in docs:
        require_approval(approvals, "gorden-generation-scope", root, root / DOCS["gorden_component"])

    if scenario == "new-deck":
        planned_ids = slide_ids(docs["slide_plan"], "slide_plan.json")
        representative_count = min(len(planned_ids), 3)
        style_reference_plan = docs["style_reference_plan"]
        style_reference_approval = require_approval(
            approvals,
            "style-reference-choice",
            root,
            root / DOCS["style_reference_plan"],
        )
        style_reference_decision = style_reference_approval.get("decision", {})
        if (
            style_reference_decision.get("selected_ids") != style_reference_plan["selected_ids"]
            or style_reference_decision.get("usage_mode") != style_reference_plan["usage_mode"]
        ):
            raise ContractError("style-reference-choice differs from style_reference_plan.json")
        layouts = artifact_by_purpose(artifacts, "layout-preview")
        if len(layouts) != 1 or layouts[0]["expected_pages"] != len(planned_ids):
            raise ContractError("new-deck requires exactly one complete layout-preview PPTX")
        require_artifact_parent(root, layouts[0], root / DOCS["slide_plan"])
        layout_approval = require_approval(
            approvals, "layout-preview", root, artifact_path(root, layouts[0])
        )
        if parse_timestamp(
            style_reference_approval["confirmed_at"], "style-reference-choice.confirmed_at"
        ) < parse_timestamp(layout_approval["confirmed_at"], "layout-preview.confirmed_at"):
            raise ContractError("style-reference-choice must occur after layout-preview approval")
        styles = []
        for purpose in ("style-option-a", "style-option-b", "style-option-c"):
            found = artifact_by_purpose(artifacts, purpose)
            if (
                len(found) != 1
                or found[0]["expected_pages"] != representative_count
                or found[0]["reusable"] is not True
            ):
                raise ContractError(
                    f"{purpose} must be one reusable {representative_count}-page PPTX"
                )
            require_artifact_parent(root, found[0], artifact_path(root, layouts[0]))
            require_artifact_parent(root, found[0], root / DOCS["style_reference_plan"])
            if len(inspect_image_deck_pptx(artifact_path(root, found[0]))) != representative_count:
                raise ContractError(
                    f"{purpose} must contain {representative_count} full-slide PNG pages"
                )
            styles.extend(found)
        choices = [
            candidate
            for candidate in approvals
            if candidate.get("gate") == "style-choice"
            and Path(candidate["subject_path"]).as_posix()
            in {Path(item["path"]).as_posix() for item in styles}
        ]
        if len(choices) != 1:
            raise ContractError("missing user style-choice approval for one style PPTX")
        representatives = choices[0].get("decision", {}).get("representatives")
        expected_roles = ["cover", "regular", "complex"][:representative_count]
        if not isinstance(representatives, list) or len(representatives) != representative_count:
            raise ContractError("style-choice must map every reusable representative page")
        representative_roles = [
            item.get("role") for item in representatives if isinstance(item, dict)
        ]
        representative_ids = [
            item.get("slide_id") for item in representatives if isinstance(item, dict)
        ]
        if (
            representative_roles != expected_roles
            or len(representative_ids) != representative_count
            or len(set(representative_ids)) != representative_count
            or any(slide_id not in planned_ids for slide_id in representative_ids)
        ):
            raise ContractError(
                "style-choice must map cover/regular/complex roles to distinct planned slides"
            )

    elif scenario == "template-fill":
        require_approval(approvals, "template-analysis", root, root / DOCS["template_analysis"])
        require_approval(approvals, "template-adaptation", root, root / DOCS["adaptation_report"])
        require_approval(approvals, "template-fill-plan", root, root / DOCS["fill_plan"])
        analysis = docs["template_analysis"]
        if analysis.get("ambiguous_texts"):
            require_approval(approvals, "template-ambiguities", root, root / DOCS["template_analysis"])
        if analysis.get("missing_fonts"):
            require_approval(approvals, "missing-fonts", root, root / DOCS["template_analysis"])
        if docs["fill_plan"].get("capacity_conflicts"):
            require_approval(approvals, "template-capacity-conflict", root, root / DOCS["fill_plan"])
        reps = artifact_by_purpose(artifacts, "template-representatives")
        if len(reps) != 1 or reps[0]["expected_pages"] < 2 or reps[0]["reusable"] is not True:
            raise ContractError("template-fill requires one reusable representative PPTX with at least 2 pages")
        require_approval(approvals, "template-representatives", root, artifact_path(root, reps[0]))

    elif scenario == "reconstruction":
        if "scene1_handoff" in docs:
            require_approval(approvals, "start-reconstruction", root, root / DOCS["scene1_handoff"])
        if docs["reconstruction_plan"].get("uncertain_text"):
            require_approval(approvals, "uncertain-text", root, root / DOCS["reconstruction_plan"])
        reconstruction_assets = docs["asset_plan"].get("items", [])
        if any(item.get("asset_type") == "png" for item in reconstruction_assets):
            require_approval(approvals, "reconstruction-png-assets", root, root / DOCS["asset_plan"])
        if any(item.get("contains_semantic_text") is True for item in reconstruction_assets):
            require_approval(approvals, "artistic-text-assets", root, root / DOCS["asset_plan"])
        reps = artifact_by_purpose(artifacts, "reconstruction-prototypes")
        if len(reps) != 1 or reps[0]["expected_pages"] != 2 or reps[0]["reusable"] is not True:
            raise ContractError("reconstruction requires one reusable 2-page prototype PPTX")
        require_approval(approvals, "reconstruction-prototypes", root, artifact_path(root, reps[0]))
        qa = load_json(root / reps[0]["qa_path"])
        diagnostics = qa.get("diagnostics")
        if not isinstance(diagnostics, list) or len(diagnostics) < reps[0]["expected_pages"] * 3:
            raise ContractError("reconstruction QA requires side-by-side, blend and diff diagnostics per page")
        for item in diagnostics:
            if not resolve_inside(root, item, "reconstruction diagnostic").is_file():
                raise ContractError(f"reconstruction diagnostic missing: {item}")

    elif scenario == "reference-deck":
        profile = docs["reference_profile"]
        if profile.get("conflicts"):
            require_approval(approvals, "primary-reference", root, root / DOCS["reference_profile"])
        layouts = artifact_by_purpose(artifacts, "layout-preview")
        reps = artifact_by_purpose(artifacts, "reference-representatives")
        if len(layouts) != 1 or len(reps) != 1 or reps[0]["expected_pages"] != 3 or reps[0]["reusable"] is not True:
            raise ContractError("reference-deck requires one layout preview and one reusable 3-page representative PPTX")
        require_approval(approvals, "layout-preview", root, artifact_path(root, layouts[0]))
        require_approval(approvals, "reference-representatives", root, artifact_path(root, reps[0]))

    elif scenario == "existing-deck":
        require_approval(approvals, "change-scope", root, root / DOCS["change_plan"])
        if docs["change_plan"].get("keep_original_slides") is False:
            require_approval(approvals, "delete-original-slides", root, root / DOCS["change_plan"])
        if project["new_page_count"] > 0:
            layouts = artifact_by_purpose(artifacts, "layout-preview")
            if len(layouts) != 1 or layouts[0]["expected_pages"] != project["new_page_count"]:
                raise ContractError("existing-deck new pages require a complete layout-preview PPTX")
            require_approval(approvals, "layout-preview", root, artifact_path(root, layouts[0]))


def validate_parents(root: Path, build: dict[str, Any]) -> None:
    parents = build.get("parents")
    if not isinstance(parents, list) or not parents:
        raise ContractError("build_plan.parents must be a non-empty array")
    for index, parent in enumerate(parents):
        require(parent, ["path", "sha256"], f"build_plan.parents[{index}]")
        path = resolve_inside(root, parent["path"], "build parent")
        if not path.is_file() or parent["sha256"] != sha256(path):
            raise ContractError(f"build parent is missing or stale: {parent['path']}")


def require_build_parent(root: Path, build: dict[str, Any], required_path: Path) -> None:
    relative = required_path.resolve().relative_to(root).as_posix()
    current_hash = sha256(required_path)
    if not any(
        Path(parent.get("path", "")).as_posix() == relative
        and parent.get("sha256") == current_hash
        for parent in build["parents"]
    ):
        raise ContractError(f"build_plan.json does not bind required parent: {relative}")


def validate_final(
    root: Path,
    project: dict[str, Any],
    docs: dict[str, dict[str, Any]],
    artifacts: list[dict[str, Any]],
    approvals: list[dict[str, Any]],
) -> None:
    build = load_doc(root, "build_plan", project["run_id"])
    validate_parents(root, build)
    scenario = project["scenario"]
    required_parents = {
        "new-deck": [
            root / DOCS["slide_plan"],
            root / DOCS["style_reference_plan"],
            root / DOCS["asset_plan"],
            root / DOCS["gorden_component"],
        ],
        "template-fill": [root / DOCS["fill_plan"]],
        "reconstruction": [root / DOCS["reconstruction_plan"], root / DOCS["asset_plan"], root / DOCS["gorden_component"]],
        "reference-deck": [root / DOCS["reference_profile"], root / DOCS["slide_plan"], root / DOCS["asset_plan"]],
        "existing-deck": [root / DOCS["change_plan"]],
    }[scenario]
    if scenario == "new-deck":
        required_parents.append(resolve_inside(root, require_approval(approvals, "style-choice", root)["subject_path"], "style choice"))
    elif scenario == "template-fill":
        required_parents.append(artifact_path(root, artifact_by_purpose(artifacts, "template-representatives")[0]))
    elif scenario == "reconstruction":
        required_parents.append(artifact_path(root, artifact_by_purpose(artifacts, "reconstruction-prototypes")[0]))
    elif scenario == "reference-deck":
        required_parents.append(artifact_path(root, artifact_by_purpose(artifacts, "reference-representatives")[0]))
    elif scenario == "existing-deck" and project["new_page_count"] > 0:
        required_parents.append(artifact_path(root, artifact_by_purpose(artifacts, "layout-preview")[0]))
    for required_parent in required_parents:
        require_build_parent(root, build, required_parent)
    if scenario == "new-deck":
        planned_ids = slide_ids(docs["slide_plan"], "slide_plan.json")
        build_ids = slide_ids(build, "build_plan.json")
        if build_ids != planned_ids:
            raise ContractError("new-deck build must preserve the confirmed slide plan roster/order")
        for slide in build["slides"]:
            objects = slide.get("objects")
            if not isinstance(objects, list) or len(objects) != 1:
                raise ContractError("new-deck final slides must contain exactly one full-slide image")
            item = objects[0]
            if (
                item.get("kind") != "raster"
                or item.get("build_method") != "full_slide_image"
                or item.get("editable") is not False
                or item.get("bbox") != [0, 0, 1, 1]
            ):
                raise ContractError("new-deck final object must be a non-editable full-slide image")
            planned_assets = {
                asset["pages"][0]: asset
                for asset in docs["asset_plan"].get("items", [])
                if asset.get("visual_role") == "full-slide-image"
            }
            planned_asset = planned_assets[slide["id"]]
            source_path = resolve_inside(root, planned_asset["source"], "new-deck final PNG")
            if (
                item.get("asset_id") != planned_asset["id"]
                or Path(str(item.get("source_path", ""))).as_posix() != Path(planned_asset["source"]).as_posix()
                or item.get("source_sha256") != sha256(source_path)
            ):
                raise ContractError("new-deck final image must reference its confirmed full-slide asset")
    reconstruction_png_placements: set[tuple[str, str]] = set()
    if scenario == "reconstruction":
        planned_assets = {item["id"]: item for item in docs["asset_plan"].get("items", [])}
        used_asset_ids: set[str] = set()
        for slide in build.get("slides", []):
            slide_id = slide.get("id")
            for item in slide.get("objects", []):
                if item.get("kind") not in {"native_text", "svg"}:
                    if item.get("kind") != "raster":
                        raise ContractError("reconstruction build may contain only native_text, svg and planned PNG objects")
                if item.get("kind") in {"svg", "raster"}:
                    asset_id = item.get("asset_id")
                    if not isinstance(asset_id, str) or asset_id not in planned_assets:
                        raise ContractError("reconstruction visual objects must reference a planned asset_id")
                    planned = planned_assets[asset_id]
                    expected_type = "svg" if item["kind"] == "svg" else "png"
                    if planned.get("asset_type") != expected_type:
                        raise ContractError("reconstruction build object kind does not match its planned asset type")
                    if slide_id not in planned.get("pages", []):
                        raise ContractError("reconstruction build uses an asset on an unplanned page")
                    if item["kind"] == "raster" and item.get("editable") is not False:
                        raise ContractError("reconstruction PNG objects must be reported as non-editable")
                    used_asset_ids.add(asset_id)
                    if expected_type == "png":
                        reconstruction_png_placements.add((slide_id, asset_id))
        if used_asset_ids != set(planned_assets):
            raise ContractError("reconstruction build assets must exactly match the confirmed asset plan")

    finals = artifact_by_purpose(artifacts, "final-output")
    if len(finals) != 1:
        raise ContractError("final phase requires exactly one final-output PPTX")
    final = finals[0]
    if scenario == "new-deck":
        if final["expected_pages"] != len(docs["slide_plan"]["slides"]):
            raise ContractError("new-deck final image PPTX must cover every planned slide")
        final_qa = load_json(resolve_inside(root, final["qa_path"], "new-deck final QA"))
        checks = final_qa.get("checks", {})
        if checks.get("semantic_text_exact") != "pass" or checks.get("slide_plan_match") != "pass":
            raise ContractError("new-deck final QA must pass exact visible text and slide-plan checks")
        planned_assets = {
            asset["pages"][0]: asset
            for asset in docs["asset_plan"].get("items", [])
            if asset.get("visual_role") == "full-slide-image"
        }
        validate_image_deck_pptx(
            artifact_path(root, final),
            root,
            [planned_assets[slide_id] for slide_id in slide_ids(docs["slide_plan"], "slide_plan.json")],
        )
        style_choice = require_approval(approvals, "style-choice", root)
        selected_style = resolve_inside(root, style_choice["subject_path"], "selected style PPTX")
        selected_hashes = inspect_image_deck_pptx(selected_style)
        representatives = style_choice.get("decision", {}).get("representatives", [])
        for index, representative in enumerate(representatives):
            slide_id = representative["slide_id"]
            source = resolve_inside(root, planned_assets[slide_id]["source"], "reused style page PNG")
            if selected_hashes[index] != sha256(source):
                raise ContractError(f"selected style page {index + 1} was not reused as final slide {slide_id}")
        require_approval(approvals, "image-deck-final", root, artifact_path(root, final))
    build_path = root / DOCS["build_plan"]
    if not any(
        Path(parent.get("path", "")).as_posix() == Path(DOCS["build_plan"]).as_posix()
        and parent.get("sha256") == sha256(build_path)
        for parent in final["parents"]
    ):
        raise ContractError("final-output artifact must bind the current build_plan.json")
    report = load_doc(root, "report", project["run_id"])
    if Path(report.get("output_pptx", "")).as_posix() != Path(final["path"]).as_posix():
        raise ContractError("editability report does not point to final-output PPTX")
    if report.get("output_sha256") != final["sha256"]:
        raise ContractError("editability report output hash mismatch")
    slides = report.get("slides")
    if not isinstance(slides, list) or len(slides) != final["expected_pages"]:
        raise ContractError("editability report must cover every final slide")
    if scenario == "new-deck":
        planned_assets = {
            asset["pages"][0]: asset
            for asset in docs["asset_plan"].get("items", [])
            if asset.get("visual_role") == "full-slide-image"
        }
        if slide_ids(report, "editability_report.json") != list(planned_assets):
            raise ContractError("new-deck editability report must preserve the confirmed slide roster/order")
        for index, slide in enumerate(slides):
            require(
                slide,
                ["id", "native_text", "native_objects", "svg_assets", "raster_assets", "warnings"],
                f"editability_report.slides[{index}]",
            )
            rasters = slide.get("raster_assets")
            if not isinstance(slide["warnings"], list):
                raise ContractError(f"new-deck editability report slide {index + 1} warnings must be an array")
            if (
                slide.get("native_text") != 0
                or slide.get("native_objects") != 0
                or slide.get("svg_assets") != 0
                or not isinstance(rasters, list)
                or len(rasters) != 1
                or rasters[0].get("asset_id") != planned_assets[slide.get("id")]["id"]
                or rasters[0].get("visual_role") != "full-slide-image"
                or rasters[0].get("editable") is not False
                or rasters[0].get("contains_semantic_text") is not True
            ):
                raise ContractError(f"new-deck editability report must disclose slide {index + 1} as one full-slide image")
            require(
                rasters[0],
                [
                    "asset_id", "format", "visual_role", "source_kind", "editable",
                    "contains_semantic_text", "editability_impact",
                ],
                f"editability_report.slides[{index}].raster_assets[0]",
            )
            if (
                rasters[0]["format"] != "png"
                or rasters[0]["source_kind"] != "generated"
                or not str(rasters[0]["editability_impact"]).strip()
            ):
                raise ContractError(f"new-deck slide {index + 1} raster disclosure is incomplete")
        limitations = report.get("limitations")
        if not isinstance(limitations, list) or not limitations or any(not str(item).strip() for item in limitations):
            raise ContractError("new-deck report must disclose that image content is not editable")
    if scenario == "reconstruction":
        planned_assets = {item["id"]: item for item in docs["asset_plan"].get("items", [])}
        reported_png_placements: set[tuple[str, str]] = set()
        for slide in slides:
            raster_assets = slide.get("raster_assets")
            if not isinstance(raster_assets, list):
                raise ContractError("reconstruction raster_assets must be an array")
            for index, item in enumerate(raster_assets):
                context = f"editability_report {slide.get('id')}.raster_assets[{index}]"
                if not isinstance(item, dict):
                    raise ContractError(f"{context} must be an object")
                require(
                    item,
                    [
                        "asset_id", "format", "visual_role", "source_kind", "editable",
                        "contains_semantic_text", "editability_impact",
                    ],
                    context,
                )
                asset_id = item["asset_id"]
                planned = planned_assets.get(asset_id)
                if planned is None or planned.get("asset_type") != "png":
                    raise ContractError("editability report contains an unplanned raster asset")
                if slide.get("id") not in planned.get("pages", []):
                    raise ContractError("editability report assigns a PNG to an unplanned page")
                if item["format"] != "png" or item["editable"] is not False:
                    raise ContractError("reconstruction raster assets must be disclosed as non-editable PNG")
                for field in ("visual_role", "source_kind", "contains_semantic_text", "editability_impact"):
                    if item[field] != planned[field]:
                        raise ContractError(f"editability report PNG {field} differs from the confirmed asset plan")
                reported_png_placements.add((slide.get("id"), asset_id))
        if reported_png_placements != reconstruction_png_placements:
            raise ContractError("editability report PNG assets must exactly match the reconstruction build")
        qa = load_json(root / final["qa_path"])
        diagnostics = qa.get("diagnostics")
        if not isinstance(diagnostics, list) or len(diagnostics) < final["expected_pages"] * 3:
            raise ContractError("final reconstruction QA requires three diagnostics per page")
    if scenario == "existing-deck":
        change_map = report.get("change_map")
        if not isinstance(change_map, list) or not change_map:
            raise ContractError("existing-deck report must include change_map")
        if docs["change_plan"].get("keep_original_slides") is True:
            for index, mapping in enumerate(change_map):
                require(
                    mapping,
                    ["source_slide", "output_original_slide", "modified_copy_slide"],
                    f"editability_report.change_map[{index}]",
                )
                if mapping["modified_copy_slide"] != mapping["output_original_slide"] + 1:
                    raise ContractError("modified copy must immediately follow its retained original slide")
            final_qa = load_json(root / final["qa_path"])
            if final_qa.get("checks", {}).get("source_page_unchanged") != "pass":
                raise ContractError("existing-deck QA must prove retained source pages are unchanged")


def validate_run(root: Path, phase: str) -> None:
    root = root.expanduser().resolve()
    if phase not in PHASES:
        raise ContractError(f"unsupported phase: {phase}")
    project = validate_project_file(root)
    docs = validate_plan_docs(root, project, phase)
    approvals = approval_records(root)
    _manifest, artifacts = validate_artifacts(root, project["run_id"])
    validate_gorden_generation_manifest(root, project, docs, artifacts, approvals, phase)
    if phase in {"preview", "final"}:
        require_preview_contract(root, project, docs, artifacts, approvals)
    if phase == "final":
        validate_final(root, project, docs, artifacts, approvals)


def write_json(path: Path, value: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def write_dummy_pptx(path: Path, pages: int, full_slide_images: list[Path] | None = None) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if full_slide_images is not None:
        if len(full_slide_images) != pages:
            raise ValueError("full_slide_images must match pages")
        from pptx import Presentation

        presentation = Presentation()
        blank = presentation.slide_layouts[6]
        for image_path in full_slide_images:
            slide = presentation.slides.add_slide(blank)
            slide.shapes.add_picture(
                str(image_path),
                0,
                0,
                width=presentation.slide_width,
                height=presentation.slide_height,
            )
        presentation.save(path)
        return
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        for index in range(1, pages + 1):
            archive.writestr(f"ppt/slides/slide{index}.xml", "<p:sld xmlns:p='p'/>")


def write_test_png(path: Path, width: int, height: int, gray: int = 0) -> None:
    def chunk(chunk_type: bytes, payload: bytes) -> bytes:
        crc = binascii.crc32(chunk_type + payload) & 0xFFFFFFFF
        return struct.pack(">I", len(payload)) + chunk_type + payload + struct.pack(">I", crc)

    path.parent.mkdir(parents=True, exist_ok=True)
    ihdr = struct.pack(">IIBBBBB", width, height, 8, 0, 0, 0, 0)
    if not 0 <= gray <= 255:
        raise ValueError("gray must be between 0 and 255")
    scanlines = (b"\x00" + bytes([gray]) * width) * height
    path.write_bytes(
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", ihdr)
        + chunk(b"IDAT", zlib.compress(scanlines, 9))
        + chunk(b"IEND", b"")
    )


def base_doc(version: str, run_id: str, status: str = "confirmed") -> dict[str, Any]:
    return {"schema_version": version, "run_id": run_id, "status": status}


def add_artifact(
    root: Path,
    artifacts: list[dict[str, Any]],
    artifact_id: str,
    purpose: str,
    pages: int,
    reusable: bool,
    parent_paths: list[Path],
    reconstruction: bool = False,
    full_slide_images: list[Path] | None = None,
) -> dict[str, Any]:
    pptx_rel = f"previews/{artifact_id}.pptx" if purpose != "final-output" else f"exports/{artifact_id}.pptx"
    pptx = root / pptx_rel
    write_dummy_pptx(pptx, pages, full_slide_images)
    render_rel = f"qa/{artifact_id}/render"
    render_dir = root / render_rel
    render_dir.mkdir(parents=True, exist_ok=True)
    rendered: list[str] = []
    for index in range(1, pages + 1):
        image_rel = f"{render_rel}/slide-{index}.png"
        write_test_png(root / image_rel, 1600, 900)
        rendered.append(image_rel)
    diagnostics: list[str] = []
    if reconstruction:
        for index in range(1, pages + 1):
            for name in ("side-by-side", "blend", "diff-heatmap"):
                rel = f"qa/{artifact_id}/{name}-{index}.png"
                write_test_png(root / rel, 1600, 900)
                diagnostics.append(rel)
    qa_rel = f"qa/{artifact_id}/visual_qa.json"
    qa = {
        "schema_version": "pptx-workshop.visual-qa.v1",
        "artifact_id": artifact_id,
        "pptx_path": pptx_rel,
        "pptx_sha256": sha256(pptx),
        "render_dir": render_rel,
        "rendered_pages": rendered,
        "inspected_pages": list(range(1, pages + 1)),
        "inspected_by": "agent",
        "status": "pass",
        "checks": {"overflow": "pass", "wrapping": "pass", "layering": "pass"},
        "issues": [],
        "diagnostics": diagnostics,
    }
    write_json(root / qa_rel, qa)
    artifact = {
        "id": artifact_id,
        "purpose": purpose,
        "path": pptx_rel,
        "sha256": sha256(pptx),
        "qa_path": qa_rel,
        "expected_pages": pages,
        "reusable": reusable,
        "parents": [
            {"path": parent.relative_to(root).as_posix(), "sha256": sha256(parent)}
            for parent in parent_paths
        ],
    }
    artifacts.append(artifact)
    return artifact


def approve(root: Path, gate: str, subject: Path, decision: dict[str, Any] | None = None) -> None:
    rel = subject.resolve().relative_to(root.resolve()).as_posix()
    write_json(
        root / "approvals" / f"{gate}.json",
        {
            "schema_version": "pptx-workshop.approval.v1",
            "gate": gate,
            "status": "confirmed",
            "confirmed_by": "user",
            "subject_path": rel,
            "subject_sha256": sha256(subject),
            "confirmed_at": datetime.now(timezone.utc).isoformat(),
            "decision": decision or {},
        },
    )


def write_gorden_fixture(root: Path, run_id: str, scenario: str, page_count: int) -> None:
    commit = "8c05583dab8334182b71738e8dfbbec5c56a1951"
    component_root = Path(__file__).resolve().parent.parent / "components" / "gorden"
    probe_path = root / "manifests" / "backend_probe.json"
    write_json(
        probe_path,
        {
            "schema": "pptx-workshop.backend-probe.v1",
            "gorden_super_ppt": {
                "root": str(component_root.resolve()),
                "bundled": True,
                "available": True,
                "source_sha": commit,
                "image_ppt_gen": {"available": True, "compose_pptx": True},
                "image_to_pptx": {
                    "available": True,
                    "compose_pptx": True,
                    "layout_guard": True,
                    "placement_qa": True,
                    "visual_compare_qa": True,
                },
                "standard_license_file": False,
                "attribution_notice": True,
                "imagegen_required": True,
            },
            "python_modules": {"python_pptx": True, "Pillow": True, "numpy": True},
        },
    )
    pipeline = "image-ppt" if scenario == "new-deck" else "image2pptx"
    stages = []
    if scenario == "new-deck":
        stages.append(
            {
                "component": "GordenImagePPTGen",
                "purpose": "create visual source slides from the approved layout and selected style",
                "output_root": "work/gorden/image-deck",
                "manifest_path": "manifests/gorden-generation.json",
            }
        )
    else:
        stages.append(
            {
                "component": "GordenImage2PPTX",
                "purpose": "reconstruct native ordinary text over fidelity-first visual layers",
                "output_root": "work/gorden/editable",
                "manifest_path": "manifests/gorden-generation.json",
            }
        )
    representative_count = min(page_count, 3)
    minimum_calls = page_count + 2 * representative_count if scenario == "new-deck" else page_count * 3
    component_doc = base_doc(SCHEMA_VERSIONS["gorden_component"], run_id)
    component_doc.update(
        {
            "repository": "GordenSun/GordenSuperPPTSkills",
            "expected_commit": commit,
            "component_location": "components/gorden",
            "notice_path": "components/gorden/NOTICE.md",
            "probe_report": probe_path.relative_to(root).as_posix(),
            "probe_sha256": sha256(probe_path),
            "integration_mode": "bundled-attributed",
            "pipeline": pipeline,
            "stages": stages,
            "generation_scope": {
                "page_count": page_count,
                "estimated_imagegen_calls_min": minimum_calls,
                "basis": (
                    f"three {representative_count}-page style options, reuse the selected pages, then generate remaining final pages"
                    if scenario == "new-deck"
                    else "minimum 3 imagegen layer calls per source page"
                ),
                "user_notice": "fixture notice",
            },
            "source_bundled": True,
            "attribution_url": "https://github.com/GordenSun/GordenSuperPPTSkills",
        }
    )
    write_json(root / DOCS["gorden_component"], component_doc)


def make_fixture(root: Path, scenario: str, page_count: int = 1) -> None:
    if page_count < 1 or (scenario != "new-deck" and page_count != 1):
        raise ValueError("page_count customization is supported only for new-deck fixtures")
    run_id = f"test-{scenario}"
    source = root / "external-source.bin"
    source.write_bytes(b"source")
    new_pages = 1 if scenario == "existing-deck" else 0
    write_json(
        root / "project.json",
        {
            "schema_version": "pptx-workshop.project.v1",
            "run_id": run_id,
            "scenario": scenario,
            "phase": "final",
            "source_files": [{"path": str(source.resolve()), "sha256": sha256(source), "role": "primary"}],
            "new_page_count": new_pages,
        },
    )
    brief = base_doc(SCHEMA_VERSIONS["brief"], run_id)
    brief.update(
        {
            "scenario": scenario,
            "audience": "测试受众",
            "objective": "验证通路",
            "success_criteria": ["合同通过"],
            "source_policy": {},
            "constraints": {},
        }
    )
    write_json(root / DOCS["brief"], brief)

    outline = base_doc(SCHEMA_VERSIONS["outline"], run_id)
    planned_page_count = page_count if scenario == "new-deck" else 1
    outline["slides"] = [
        {
            "id": f"P{index:02d}",
            "title": f"测试 {index}",
            "purpose": "验证",
            "key_message": "通过",
            "audience_move": "理解",
            "evidence_ids": [],
        }
        for index in range(1, planned_page_count + 1)
    ]
    slide_plan = base_doc(SCHEMA_VERSIONS["slide_plan"], run_id)
    slide_plan["slides"] = [
        {
            "id": f"P{index:02d}",
            "text_locked": True,
            "layout_locked": True,
            "content_blocks": [
                {
                    "id": "B01",
                    "type": "title",
                    "text": f"测试 {index}",
                    "bbox": [0.05, 0.08, 0.9, 0.12],
                    "reading_order": 1,
                    "align": "left",
                }
            ],
            "layout": {
                "pattern": "hero",
                "reading_order": ["B01"],
                "decorative_safe_zones": [
                    {
                        "id": "decor-01",
                        "bbox": [0.05, 0.25, 0.9, 0.65],
                        "purpose": "style-only decoration and visual assets",
                    }
                ],
            },
            "asset_ids": [],
            "native_objects": ["native_text"],
            "qa_assertions": ["文字可编辑"],
        }
        for index in range(1, planned_page_count + 1)
    ]
    evidence = base_doc(SCHEMA_VERSIONS["evidence_plan"], run_id)
    evidence["items"] = []
    asset = base_doc(SCHEMA_VERSIONS["asset_plan"], run_id)
    asset["items"] = []
    style_reference_plan: dict[str, Any] | None = None
    style_reference_ids: list[str] = []

    if scenario == "new-deck":
        catalog = load_json(STYLE_LIBRARY_CATALOG)
        candidate_entries = catalog["entries"][:3]
        candidates = [
            {
                "id": entry["id"],
                "source_type": "bundled",
                "catalog_entry_id": entry["id"],
                "style_label": entry["style_name"],
                "fit_reason": "fixture candidate shown after the layout preview",
                "limitations": entry.get("limitations", []),
            }
            for entry in candidate_entries
        ]
        style_reference_ids = [candidates[0]["id"]]
        style_reference_plan = base_doc(SCHEMA_VERSIONS["style_reference_plan"], run_id)
        style_reference_plan.update(
            {
                "catalog_schema_version": STYLE_LIBRARY_VERSION,
                "catalog_sha256": sha256(STYLE_LIBRARY_CATALOG),
                "candidates": candidates,
                "selected_ids": style_reference_ids,
                "usage_mode": "direct",
                "user_direction": "Use the selected reference after preserving the confirmed layout.",
            }
        )
        asset["items"] = []
        for index, slide in enumerate(slide_plan["slides"], 1):
            slide_id = slide["id"]
            asset_id = f"slide-image-{slide_id}"
            slide["asset_ids"] = [asset_id]
            slide["native_objects"] = []
            slide["qa_assertions"] = ["最终可见文字与策划逐字一致"]
            page_image = f"assets/slide-{slide_id}.png"
            write_test_png(root / page_image, 1920, 1080, index * 10)
            asset["items"].append(
                {
                    "id": asset_id,
                    "pages": [slide_id],
                    "purpose": "final generated image for the complete slide",
                    "asset_type": "png",
                    "source": page_image,
                    "contains_semantic_text": True,
                    "ordinary_text_in_asset": True,
                    "text_removal_status": "not-applicable",
                    "visual_role": "full-slide-image",
                    "source_kind": "generated",
                    "selection_reason": "scene 1 intentionally delivers a fidelity-first image deck",
                    "editability_impact": "all visible content is flattened into the slide image",
                    "bbox": [0, 0, 1, 1],
                    "pixel_width": 1920,
                    "pixel_height": 1080,
                }
            )

    if scenario in {"new-deck", "reference-deck"}:
        write_json(root / DOCS["outline"], outline)
        write_json(root / DOCS["slide_plan"], slide_plan)
        write_json(root / DOCS["evidence_plan"], evidence)
        write_json(root / DOCS["asset_plan"], asset)
        if style_reference_plan is not None:
            write_json(root / DOCS["style_reference_plan"], style_reference_plan)
    if scenario == "template-fill":
        analysis = base_doc(SCHEMA_VERSIONS["template_analysis"], run_id)
        analysis.update({"template_kind": "framework-only", "preserve": ["Theme", "Master", "Layout"], "slide_families": [{"id": "family-1"}], "ambiguous_texts": [], "missing_fonts": []})
        adaptation = base_doc(SCHEMA_VERSIONS["adaptation_report"], run_id)
        adaptation.update({"content_regions": [{}], "capacity_risks": [], "unsupported_objects": []})
        fill = base_doc(SCHEMA_VERSIONS["fill_plan"], run_id)
        fill.update({"slides": [{"id": "P01", "source_slide": 1, "layout_rationale": "容量匹配", "edits": []}], "capacity_conflicts": []})
        write_json(root / DOCS["template_analysis"], analysis)
        write_json(root / DOCS["adaptation_report"], adaptation)
        write_json(root / DOCS["fill_plan"], fill)
    if scenario == "reconstruction":
        reconstruction = base_doc(SCHEMA_VERSIONS["reconstruction_plan"], run_id)
        reconstruction.update({"source_locked": True, "visual_asset_policy": "fidelity-first-controlled-png", "slides": [{"id": "P01", "source_page": 1, "ref_width": 1920, "ref_height": 1080, "objects": [{"id": "t1", "semantic_role": "text", "target_kind": "native_text"}]}], "uncertain_text": []})
        frame_rel = "assets/frame-p01.png"
        write_test_png(root / frame_rel, 1920, 1080)
        asset["items"] = [
            {
                "id": "frame-p01",
                "pages": ["P01"],
                "purpose": "preserve the complex visual frame",
                "asset_type": "png",
                "source": frame_rel,
                "contains_semantic_text": False,
                "ordinary_text_in_asset": False,
                "semantic_text_policy": "none",
                "text_removal_status": "verified-clean",
                "visual_role": "frame",
                "source_kind": "extracted",
                "selection_reason": "native or SVG redraw would reduce visual fidelity",
                "editability_impact": "the frame is replaceable but not object-level editable",
                "bbox": [0, 0, 1, 1],
                "pixel_width": 1920,
                "pixel_height": 1080,
            }
        ]
        write_json(root / DOCS["reconstruction_plan"], reconstruction)
        write_json(root / DOCS["asset_plan"], asset)
    if scenario == "reference-deck":
        profile = base_doc(SCHEMA_VERSIONS["reference_profile"], run_id)
        profile.update({"primary_reference": "primary", "references": [{"id": "primary"}], "design_rules": {}, "conflicts": []})
        write_json(root / DOCS["reference_profile"], profile)
    if scenario == "existing-deck":
        write_json(root / DOCS["outline"], outline)
        write_json(root / DOCS["slide_plan"], slide_plan)
        change = base_doc(SCHEMA_VERSIONS["change_plan"], run_id)
        change.update({"source_pptx": str(source.resolve()), "changes": [{"source_slide": 1, "selector": "title", "before": "旧", "after": "新", "preserve": ["font", "font_size", "color", "position", "size", "paragraph", "fill", "border"]}], "new_pages": [{"id": "P01"}], "keep_original_slides": True})
        write_json(root / DOCS["change_plan"], change)

    if scenario in {"new-deck", "reconstruction"}:
        write_gorden_fixture(root, run_id, scenario, planned_page_count)

    artifacts: list[dict[str, Any]] = []
    if scenario == "new-deck":
        layout = add_artifact(
            root,
            artifacts,
            "layout",
            "layout-preview",
            planned_page_count,
            False,
            [root / DOCS["slide_plan"]],
        )
        planned_ids = [slide["id"] for slide in slide_plan["slides"]]
        representative_count = min(planned_page_count, 3)
        representative_ids = (
            planned_ids
            if planned_page_count <= 3
            else [planned_ids[0], planned_ids[1], planned_ids[-1]]
        )
        representative_roles = ["cover", "regular", "complex"][:representative_count]
        style_images: dict[str, list[Path]] = {}
        for letter_index, letter in enumerate(("a", "b", "c")):
            images: list[Path] = []
            for page_index, representative_id in enumerate(representative_ids, 1):
                if letter == "a":
                    planned_asset = next(
                        item for item in asset["items"] if item["pages"] == [representative_id]
                    )
                    image_path = root / planned_asset["source"]
                else:
                    image_path = root / "assets" / "styles" / f"style-{letter}-{page_index}.png"
                    write_test_png(image_path, 1920, 1080, 20 + letter_index * 80 + page_index * 10)
                images.append(image_path)
            style_images[letter] = images
        styles = [
            add_artifact(
                root,
                artifacts,
                f"style-{letter}",
                f"style-option-{letter}",
                representative_count,
                True,
                [root / layout["path"], root / DOCS["style_reference_plan"]],
                full_slide_images=style_images[letter],
            )
            for letter in ("a", "b", "c")
        ]
        approve(root, "layout-preview", root / layout["path"])
        approve(
            root,
            "style-reference-choice",
            root / DOCS["style_reference_plan"],
            {"selected_ids": style_reference_ids, "usage_mode": "direct"},
        )
        approve(root, "asset-plan", root / DOCS["asset_plan"])
        approve(root, "gorden-generation-scope", root / DOCS["gorden_component"])
        approve(
            root,
            "style-choice",
            root / styles[0]["path"],
            {
                "selected": "style-a",
                "representatives": [
                    {"role": role, "slide_id": slide_id}
                    for role, slide_id in zip(representative_roles, representative_ids)
                ],
            },
        )
    elif scenario == "template-fill":
        reps = add_artifact(root, artifacts, "template-reps", "template-representatives", 2, True, [root / DOCS["fill_plan"]])
        approve(root, "template-analysis", root / DOCS["template_analysis"])
        approve(root, "template-adaptation", root / DOCS["adaptation_report"])
        approve(root, "template-fill-plan", root / DOCS["fill_plan"])
        approve(root, "template-representatives", root / reps["path"])
    elif scenario == "reconstruction":
        reps = add_artifact(root, artifacts, "rebuild-reps", "reconstruction-prototypes", 2, True, [root / DOCS["reconstruction_plan"]], True)
        approve(root, "asset-plan", root / DOCS["asset_plan"])
        approve(root, "reconstruction-png-assets", root / DOCS["asset_plan"])
        approve(root, "gorden-generation-scope", root / DOCS["gorden_component"])
        approve(root, "reconstruction-prototypes", root / reps["path"])
    elif scenario == "reference-deck":
        layout = add_artifact(root, artifacts, "layout", "layout-preview", 1, False, [root / DOCS["slide_plan"]])
        reps = add_artifact(root, artifacts, "reference-reps", "reference-representatives", 3, True, [root / DOCS["reference_profile"], root / layout["path"]])
        approve(root, "layout-preview", root / layout["path"])
        approve(root, "reference-representatives", root / reps["path"])
        approve(root, "asset-plan", root / DOCS["asset_plan"])
    else:
        layout = add_artifact(root, artifacts, "layout", "layout-preview", 1, False, [root / DOCS["slide_plan"]])
        approve(root, "change-scope", root / DOCS["change_plan"])
        approve(root, "layout-preview", root / layout["path"])
    parent_key = {
        "new-deck": "slide_plan",
        "template-fill": "fill_plan",
        "reconstruction": "reconstruction_plan",
        "reference-deck": "slide_plan",
        "existing-deck": "change_plan",
    }[scenario]
    parent_path = root / DOCS[parent_key]
    build_parent_paths = [parent_path]
    if scenario in {"new-deck", "reconstruction", "reference-deck"}:
        build_parent_paths.append(root / DOCS["asset_plan"])
    if scenario in {"new-deck", "reconstruction"}:
        build_parent_paths.append(root / DOCS["gorden_component"])
    if scenario == "new-deck":
        build_parent_paths.append(root / DOCS["style_reference_plan"])
        build_parent_paths.append(root / styles[0]["path"])
    elif scenario == "template-fill":
        build_parent_paths.append(root / reps["path"])
    elif scenario == "reconstruction":
        build_parent_paths.append(root / reps["path"])
    elif scenario == "reference-deck":
        build_parent_paths.extend([root / DOCS["reference_profile"], root / reps["path"]])
    elif scenario == "existing-deck" and new_pages > 0:
        build_parent_paths.append(root / layout["path"])
    if scenario == "reconstruction":
        build_slides = [
            {
                "id": "P01",
                "objects": [
                    {"id": "o1", "kind": "native_text", "bbox": [0.1, 0.1, 0.3, 0.1], "build_method": "native_text", "editable": True},
                    {"id": "o2", "kind": "raster", "asset_id": "frame-p01", "bbox": [0, 0, 1, 1], "build_method": "png_asset", "editable": False},
                ],
            }
        ]
    elif scenario == "new-deck":
        build_slides = [
            {
                "id": item["pages"][0],
                "objects": [
                    {
                        "id": item["id"],
                        "kind": "raster",
                        "asset_id": item["id"],
                        "source_path": item["source"],
                        "source_sha256": sha256(root / item["source"]),
                        "bbox": [0, 0, 1, 1],
                        "build_method": "full_slide_image",
                        "editable": False,
                    }
                ],
            }
            for item in asset["items"]
        ]
    else:
        build_slides = [
            {"id": "P01", "objects": [{"id": f"o{i}", "kind": kind, "bbox": [0.1, 0.1, 0.3, 0.1], "build_method": kind, "editable": True} for i, kind in enumerate(["native_text", "native_shape"], 1)]}
        ]
    if scenario == "existing-deck":
        build_slides.append({"id": "P02", "objects": build_slides[0]["objects"]})
    build = {
        "schema_version": SCHEMA_VERSIONS["build_plan"],
        "run_id": run_id,
        "parents": [
            {"path": item.relative_to(root).as_posix(), "sha256": sha256(item)}
            for item in build_parent_paths
        ],
        "slides": build_slides,
    }
    write_json(root / DOCS["build_plan"], build)
    final_pages = 2 if scenario == "existing-deck" else planned_page_count
    final = add_artifact(
        root,
        artifacts,
        "final",
        "final-output",
        final_pages,
        True,
        [root / DOCS["build_plan"]],
        scenario == "reconstruction",
        (
            [root / item["source"] for item in asset["items"]]
            if scenario == "new-deck"
            else None
        ),
    )
    if scenario == "existing-deck":
        final_qa = load_json(root / final["qa_path"])
        final_qa["checks"]["source_page_unchanged"] = "pass"
        write_json(root / final["qa_path"], final_qa)
    if scenario == "new-deck":
        final_qa = load_json(root / final["qa_path"])
        final_qa["checks"].update({"semantic_text_exact": "pass", "slide_plan_match": "pass"})
        write_json(root / final["qa_path"], final_qa)
        approve(root, "image-deck-final", root / final["path"])
    if scenario in {"new-deck", "reconstruction"}:
        calls: list[dict[str, Any]] = []
        qa_results: list[dict[str, Any]] = []
        if scenario == "new-deck":
            call_index = 0
            slide_by_id = {slide["id"]: slide for slide in slide_plan["slides"]}
            scope_confirmed_at = load_json(
                root / "approvals" / "gorden-generation-scope.json"
            )["confirmed_at"]
            style_confirmed_at = load_json(
                root / "approvals" / "style-choice.json"
            )["confirmed_at"]
            for letter in ("a", "b", "c"):
                for representative_id, image_path in zip(representative_ids, style_images[letter]):
                    call_index += 1
                    slide = slide_by_id[representative_id]
                    prompt_contract = "\n".join(
                        f"{block['text']} bbox={json.dumps(block['bbox'], ensure_ascii=False, separators=(',', ':'))}"
                        for block in slide["content_blocks"]
                    )
                    calls.append(
                        {
                            "id": f"call-{call_index:03d}",
                            "purpose": f"style-option-{letter}",
                            "workflow_phase": "style-selection",
                            "page_id": representative_id,
                            "slide_contract_sha256": json_sha256(slide),
                            "style_reference_ids": style_reference_ids,
                            "provider": "fixture-imagegen",
                            "model": "fixture-model",
                            "prompt": f"Generate {letter} style for {representative_id}\n{prompt_contract}",
                            "generated_at": scope_confirmed_at,
                            "status": "succeeded",
                            "output_path": image_path.relative_to(root).as_posix(),
                            "output_sha256": sha256(image_path),
                        }
                    )
            for planned_asset in asset["items"]:
                page_id = planned_asset["pages"][0]
                if page_id in representative_ids:
                    continue
                call_index += 1
                output_path = root / planned_asset["source"]
                slide = slide_by_id[page_id]
                prompt_contract = "\n".join(
                    f"{block['text']} bbox={json.dumps(block['bbox'], ensure_ascii=False, separators=(',', ':'))}"
                    for block in slide["content_blocks"]
                )
                calls.append(
                    {
                        "id": f"call-{call_index:03d}",
                        "purpose": "final-page",
                        "workflow_phase": "formal-generation",
                        "page_id": page_id,
                        "slide_contract_sha256": json_sha256(slide),
                        "style_reference_ids": style_reference_ids,
                        "provider": "fixture-imagegen",
                        "model": "fixture-model",
                        "prompt": f"Generate final page {page_id}\n{prompt_contract}",
                        "generated_at": style_confirmed_at,
                        "status": "succeeded",
                        "output_path": output_path.relative_to(root).as_posix(),
                        "output_sha256": sha256(output_path),
                    }
                )
        else:
            call_index = 0
            reconstruction_ids = [slide["id"] for slide in reconstruction["slides"]]
            for page_id in reconstruction_ids:
                for purpose in ("background", "frame", "icons"):
                    call_index += 1
                    if purpose == "frame":
                        output_path = root / asset["items"][0]["source"]
                    else:
                        output_path = root / "work" / "gorden" / "editable" / page_id / f"{purpose}.png"
                        write_test_png(output_path, 1920, 1080, 40 + call_index * 10)
                    calls.append(
                        {
                            "id": f"call-{call_index:03d}",
                            "purpose": purpose,
                            "workflow_phase": "reconstruction",
                            "page_id": page_id,
                            "provider": "fixture-imagegen",
                            "model": "fixture-model",
                            "prompt": f"Generate {purpose} for {page_id}",
                            "generated_at": "2026-07-31T00:00:00Z",
                            "status": "succeeded",
                            "output_path": output_path.relative_to(root).as_posix(),
                            "output_sha256": sha256(output_path),
                        }
                    )
            for tool in ("layout_guard", "placement_qa", "visual_compare_qa"):
                report_path = root / "qa" / "gorden" / f"{tool}.json"
                write_json(report_path, {"tool": tool, "status": "pass"})
                qa_results.append(
                    {
                        "tool": tool,
                        "status": "pass",
                        "report_path": report_path.relative_to(root).as_posix(),
                        "report_sha256": sha256(report_path),
                    }
                )
        write_json(
            root / "manifests" / "gorden-generation.json",
            {
                "schema_version": GORDEN_GENERATION_MANIFEST_VERSION,
                "run_id": run_id,
                "component": (
                    "GordenImagePPTGen" if scenario == "new-deck" else "GordenImage2PPTX"
                ),
                "calls": calls,
                "qa_results": qa_results,
            },
        )
    write_json(
        root / DOCS["artifacts"],
        {"schema_version": SCHEMA_VERSIONS["artifacts"], "run_id": run_id, "artifacts": artifacts},
    )
    report_slides = [
        {"id": f"P{index:02d}", "native_text": 1, "native_objects": 0, "svg_assets": 0, "raster_assets": [], "warnings": []}
        for index in range(1, final_pages + 1)
    ]
    if scenario == "reconstruction":
        planned_png = asset["items"][0]
        report_slides[0]["raster_assets"] = [
            {
                "asset_id": planned_png["id"],
                "format": "png",
                "visual_role": planned_png["visual_role"],
                "source_kind": planned_png["source_kind"],
                "editable": False,
                "contains_semantic_text": planned_png["contains_semantic_text"],
                "editability_impact": planned_png["editability_impact"],
            }
        ]
    if scenario == "new-deck":
        for report_slide, planned_asset in zip(report_slides, asset["items"]):
            report_slide.update(
                {
                    "native_text": 0,
                    "native_objects": 0,
                    "svg_assets": 0,
                    "raster_assets": [
                        {
                            "asset_id": planned_asset["id"],
                            "format": "png",
                            "visual_role": "full-slide-image",
                            "source_kind": "generated",
                            "editable": False,
                            "contains_semantic_text": True,
                            "editability_impact": "all visible content is flattened into the slide image",
                        }
                    ],
                }
            )
    report = {
        "schema_version": SCHEMA_VERSIONS["report"],
        "run_id": run_id,
        "output_pptx": final["path"],
        "output_sha256": final["sha256"],
        "slides": report_slides,
        "limitations": (
            ["Scene 1 is an image-based deck; start a separate reconstruction run for editable text."]
            if scenario == "new-deck"
            else []
        ),
    }
    if scenario == "existing-deck":
        report["change_map"] = [{"source_slide": 1, "output_original_slide": 1, "modified_copy_slide": 2}]
    write_json(root / DOCS["report"], report)


def attach_scene1_handoff(root: Path, upstream_root: Path) -> Path:
    project_path = root / "project.json"
    project = load_json(project_path)
    upstream_project_path = upstream_root / "project.json"
    upstream_project = load_json(upstream_project_path)
    upstream_asset_plan_path = upstream_root / DOCS["asset_plan"]
    upstream_assets = [
        item
        for item in load_json(upstream_asset_plan_path)["items"]
        if item.get("visual_role") == "full-slide-image"
    ]
    upstream_manifest_path = upstream_root / DOCS["artifacts"]
    upstream_manifest = load_json(upstream_manifest_path)
    upstream_final = next(
        item for item in upstream_manifest["artifacts"] if item["purpose"] == "final-output"
    )
    upstream_final_path = upstream_root / upstream_final["path"]
    upstream_approval_path = upstream_root / "approvals" / "image-deck-final.json"
    pages = [
        {
            "id": item["pages"][0],
            "path": str((upstream_root / item["source"]).resolve()),
            "sha256": sha256(upstream_root / item["source"]),
        }
        for item in upstream_assets
    ]
    project["source_files"] = [
        {"path": page["path"], "sha256": page["sha256"], "role": "scene1-handoff"}
        for page in pages
    ]
    write_json(project_path, project)
    handoff = base_doc(SCHEMA_VERSIONS["scene1_handoff"], project["run_id"])
    handoff.update(
        {
            "upstream_run_id": upstream_project["run_id"],
            "upstream_project": {
                "path": str(upstream_project_path.resolve()),
                "sha256": sha256(upstream_project_path),
            },
            "upstream_artifact_manifest": {
                "path": str(upstream_manifest_path.resolve()),
                "sha256": sha256(upstream_manifest_path),
            },
            "upstream_asset_plan": {
                "path": str(upstream_asset_plan_path.resolve()),
                "sha256": sha256(upstream_asset_plan_path),
            },
            "upstream_final_pptx": {
                "path": str(upstream_final_path.resolve()),
                "sha256": sha256(upstream_final_path),
            },
            "upstream_image_deck_approval": {
                "path": str(upstream_approval_path.resolve()),
                "sha256": sha256(upstream_approval_path),
            },
            "pages": pages,
        }
    )
    handoff_path = root / DOCS["scene1_handoff"]
    write_json(handoff_path, handoff)
    approve(root, "start-reconstruction", handoff_path)
    return handoff_path


def self_test() -> None:
    def expect_contract_error(action: Any, expected: str, label: str) -> None:
        try:
            action()
        except ContractError as exc:
            if expected not in str(exc):
                raise AssertionError(f"{label} failed for the wrong reason: {exc}") from exc
        else:
            raise AssertionError(label)

    def refresh_final_bindings(root: Path) -> None:
        manifest_path = root / DOCS["artifacts"]
        manifest = load_json(manifest_path)
        final = next(item for item in manifest["artifacts"] if item["purpose"] == "final-output")
        final_path = root / final["path"]
        final["sha256"] = sha256(final_path)
        write_json(manifest_path, manifest)
        qa_path = root / final["qa_path"]
        qa = load_json(qa_path)
        qa["pptx_sha256"] = final["sha256"]
        write_json(qa_path, qa)
        approval_path = root / "approvals" / "image-deck-final.json"
        approval = load_json(approval_path)
        approval["subject_sha256"] = final["sha256"]
        write_json(approval_path, approval)
        report_path = root / DOCS["report"]
        report = load_json(report_path)
        report["output_sha256"] = final["sha256"]
        write_json(report_path, report)

    with tempfile.TemporaryDirectory(prefix="pptx-workshop-tests-") as temp:
        base = Path(temp)
        for scenario in sorted(SCENARIOS):
            root = base / scenario
            root.mkdir()
            make_fixture(root, scenario)
            validate_run(root, "plan")
            validate_run(root, "preview")
            validate_run(root, "final")

        short_deck = base / "short-new-deck"
        short_deck.mkdir()
        make_fixture(short_deck, "new-deck", 2)
        validate_run(short_deck, "plan")
        validate_run(short_deck, "preview")
        validate_run(short_deck, "final")
        short_component_path = short_deck / DOCS["gorden_component"]
        short_component = load_json(short_component_path)
        if short_component["generation_scope"]["estimated_imagegen_calls_min"] != 6:
            raise AssertionError("two-page scene 1 minimum must be 3R + (N-R) = 6")
        short_component["generation_scope"]["estimated_imagegen_calls_min"] = 5
        write_json(short_component_path, short_component)
        expect_contract_error(
            lambda: validate_run(short_deck, "plan"),
            "underestimates minimum imagegen calls",
            "short scene 1 generation scope must use R=min(N,3)",
        )

        role_mapped_deck = base / "role-mapped-new-deck"
        role_mapped_deck.mkdir()
        make_fixture(role_mapped_deck, "new-deck", 4)
        validate_run(role_mapped_deck, "preview")
        validate_run(role_mapped_deck, "final")
        role_choice = load_json(role_mapped_deck / "approvals" / "style-choice.json")
        if [item["slide_id"] for item in role_choice["decision"]["representatives"]] != [
            "P01", "P02", "P04"
        ]:
            raise AssertionError("representative roles must support planned pages beyond the first three")

        missing_style_reference_approval = base / "missing-style-reference-approval"
        missing_style_reference_approval.mkdir()
        make_fixture(missing_style_reference_approval, "new-deck")
        (missing_style_reference_approval / "approvals" / "style-reference-choice.json").unlink()
        expect_contract_error(
            lambda: validate_run(missing_style_reference_approval, "preview"),
            "missing user approval: style-reference-choice",
            "scene 1 may not generate representatives before the user confirms style references",
        )

        early_style_reference_approval = base / "early-style-reference-approval"
        early_style_reference_approval.mkdir()
        make_fixture(early_style_reference_approval, "new-deck")
        style_approval_path = (
            early_style_reference_approval / "approvals" / "style-reference-choice.json"
        )
        style_approval = load_json(style_approval_path)
        style_approval["confirmed_at"] = "2000-01-01T00:00:00Z"
        write_json(style_approval_path, style_approval)
        expect_contract_error(
            lambda: validate_run(early_style_reference_approval, "preview"),
            "style-reference-choice must occur after layout-preview approval",
            "style references may not be approved before the no-style layout preview",
        )

        unknown_style_reference = base / "unknown-style-reference"
        unknown_style_reference.mkdir()
        make_fixture(unknown_style_reference, "new-deck")
        style_plan_path = unknown_style_reference / DOCS["style_reference_plan"]
        style_plan = load_json(style_plan_path)
        style_plan["candidates"][0]["id"] = "unknown-bundled-style"
        style_plan["candidates"][0]["catalog_entry_id"] = "unknown-bundled-style"
        style_plan["selected_ids"] = ["unknown-bundled-style"]
        write_json(style_plan_path, style_plan)
        expect_contract_error(
            lambda: validate_run(unknown_style_reference, "preview"),
            "unknown bundled catalog entry",
            "unknown bundled style ids must fail before representative generation",
        )

        missing_style_parent = base / "missing-style-reference-parent"
        missing_style_parent.mkdir()
        make_fixture(missing_style_parent, "new-deck")
        artifact_manifest_path = missing_style_parent / DOCS["artifacts"]
        artifact_manifest = load_json(artifact_manifest_path)
        style_a = next(
            item for item in artifact_manifest["artifacts"] if item["purpose"] == "style-option-a"
        )
        style_a["parents"] = [
            parent
            for parent in style_a["parents"]
            if Path(parent["path"]).as_posix() != Path(DOCS["style_reference_plan"]).as_posix()
        ]
        write_json(artifact_manifest_path, artifact_manifest)
        expect_contract_error(
            lambda: validate_run(missing_style_parent, "preview"),
            "does not bind required parent: plans/style_reference_plan.json",
            "style representatives must bind the confirmed style reference plan",
        )

        wrong_generation_style_reference = base / "wrong-generation-style-reference"
        wrong_generation_style_reference.mkdir()
        make_fixture(wrong_generation_style_reference, "new-deck")
        generation_path = wrong_generation_style_reference / "manifests" / "gorden-generation.json"
        generation = load_json(generation_path)
        generation["calls"][0]["style_reference_ids"] = ["unconfirmed-style"]
        write_json(generation_path, generation)
        expect_contract_error(
            lambda: validate_run(wrong_generation_style_reference, "preview"),
            "not bound to the confirmed style references",
            "scene 1 image generation must use exactly the confirmed style references",
        )

        missing_block_bbox = load_json(role_mapped_deck / DOCS["slide_plan"])
        del missing_block_bbox["slides"][0]["content_blocks"][0]["bbox"]
        expect_contract_error(
            lambda: validate_slide_plan(missing_block_bbox),
            "missing fields: bbox",
            "slide plan must lock an approximate bbox for every visible text block",
        )

        wrong_style_phase = base / "wrong-style-phase"
        wrong_style_phase.mkdir()
        make_fixture(wrong_style_phase, "new-deck")
        generation_path = wrong_style_phase / "manifests" / "gorden-generation.json"
        generation = load_json(generation_path)
        generation["calls"][0]["workflow_phase"] = "formal-generation"
        write_json(generation_path, generation)
        expect_contract_error(
            lambda: validate_run(wrong_style_phase, "preview"),
            "separate style-selection from formal-generation",
            "style representative image calls must belong to style selection",
        )

        prompt_without_text = base / "prompt-without-locked-text"
        prompt_without_text.mkdir()
        make_fixture(prompt_without_text, "new-deck")
        generation_path = prompt_without_text / "manifests" / "gorden-generation.json"
        generation = load_json(generation_path)
        generation["calls"][0]["prompt"] = "bbox=[0.05,0.08,0.9,0.12]"
        write_json(generation_path, generation)
        expect_contract_error(
            lambda: validate_run(prompt_without_text, "preview"),
            "every locked text block and its bbox",
            "scene 1 prompt must contain the complete locked text",
        )

        regenerated_representative = base / "regenerated-representative"
        regenerated_representative.mkdir()
        make_fixture(regenerated_representative, "new-deck", 4)
        generation_path = regenerated_representative / "manifests" / "gorden-generation.json"
        generation = load_json(generation_path)
        slide = load_json(regenerated_representative / DOCS["slide_plan"])["slides"][0]
        block = slide["content_blocks"][0]
        page_asset = next(
            item
            for item in load_json(regenerated_representative / DOCS["asset_plan"])["items"]
            if item["pages"] == ["P01"]
        )
        output_path = regenerated_representative / page_asset["source"]
        generation["calls"].append(
            {
                "id": "call-regenerated-representative",
                "purpose": "final-page",
                "workflow_phase": "formal-generation",
                "page_id": "P01",
                "slide_contract_sha256": json_sha256(slide),
                "style_reference_ids": load_json(
                    regenerated_representative / DOCS["style_reference_plan"]
                )["selected_ids"],
                "provider": "fixture-imagegen",
                "model": "fixture-model",
                "prompt": (
                    f"{block['text']} bbox="
                    f"{json.dumps(block['bbox'], ensure_ascii=False, separators=(',', ':'))}"
                ),
                "generated_at": load_json(
                    regenerated_representative / "approvals" / "style-choice.json"
                )["confirmed_at"],
                "status": "succeeded",
                "output_path": output_path.relative_to(regenerated_representative).as_posix(),
                "output_sha256": sha256(output_path),
            }
        )
        write_json(generation_path, generation)
        expect_contract_error(
            lambda: validate_run(regenerated_representative, "final"),
            "may not be regenerated in formal-generation",
            "selected representative pages must be reused without formal regeneration",
        )

        invalid_role_map = base / "invalid-role-map"
        invalid_role_map.mkdir()
        make_fixture(invalid_role_map, "new-deck", 4)
        role_path = invalid_role_map / "approvals" / "style-choice.json"
        role_choice = load_json(role_path)
        role_choice["decision"]["representatives"][2]["role"] = "regular"
        write_json(role_path, role_choice)
        expect_contract_error(
            lambda: validate_run(invalid_role_map, "preview"),
            "map cover/regular/complex roles",
            "style choice without the required representative role map must fail",
        )

        incomplete_layout = base / "incomplete-layout"
        incomplete_layout.mkdir()
        make_fixture(incomplete_layout, "new-deck", 2)
        artifact_manifest_path = incomplete_layout / DOCS["artifacts"]
        artifact_manifest = load_json(artifact_manifest_path)
        layout_artifact = next(
            item for item in artifact_manifest["artifacts"] if item["purpose"] == "layout-preview"
        )
        layout_path = incomplete_layout / layout_artifact["path"]
        write_dummy_pptx(layout_path, 1)
        layout_artifact["sha256"] = sha256(layout_path)
        layout_artifact["expected_pages"] = 1
        for candidate in artifact_manifest["artifacts"]:
            for parent in candidate.get("parents", []):
                if Path(parent.get("path", "")).as_posix() == Path(layout_artifact["path"]).as_posix():
                    parent["sha256"] = layout_artifact["sha256"]
        layout_qa_path = incomplete_layout / layout_artifact["qa_path"]
        layout_qa = load_json(layout_qa_path)
        layout_qa["pptx_sha256"] = layout_artifact["sha256"]
        layout_qa["rendered_pages"] = layout_qa["rendered_pages"][:1]
        layout_qa["inspected_pages"] = [1]
        write_json(layout_qa_path, layout_qa)
        layout_approval_path = incomplete_layout / "approvals" / "layout-preview.json"
        layout_approval = load_json(layout_approval_path)
        layout_approval["subject_sha256"] = layout_artifact["sha256"]
        write_json(layout_approval_path, layout_approval)
        write_json(artifact_manifest_path, artifact_manifest)
        expect_contract_error(
            lambda: validate_run(incomplete_layout, "preview"),
            "complete layout-preview",
            "scene 1 layout preview must cover every planned slide",
        )

        wrong_layout_parent = base / "wrong-layout-parent"
        wrong_layout_parent.mkdir()
        make_fixture(wrong_layout_parent, "new-deck")
        artifact_manifest_path = wrong_layout_parent / DOCS["artifacts"]
        artifact_manifest = load_json(artifact_manifest_path)
        layout_artifact = next(
            item for item in artifact_manifest["artifacts"] if item["purpose"] == "layout-preview"
        )
        brief_path = wrong_layout_parent / DOCS["brief"]
        layout_artifact["parents"] = [
            {"path": DOCS["brief"], "sha256": sha256(brief_path)}
        ]
        write_json(artifact_manifest_path, artifact_manifest)
        expect_contract_error(
            lambda: validate_run(wrong_layout_parent, "preview"),
            "does not bind required parent: plans/slide_plan.json",
            "scene 1 layout preview must bind the current slide plan",
        )

        missing_generation_manifest = base / "missing-generation-manifest"
        missing_generation_manifest.mkdir()
        make_fixture(missing_generation_manifest, "new-deck")
        (missing_generation_manifest / "manifests" / "gorden-generation.json").unlink()
        expect_contract_error(
            lambda: validate_run(missing_generation_manifest, "preview"),
            "missing file",
            "scene 1 may not preview without a real generation manifest",
        )

        incomplete_gorden_qa = base / "incomplete-gorden-qa"
        incomplete_gorden_qa.mkdir()
        make_fixture(incomplete_gorden_qa, "reconstruction")
        generation_path = incomplete_gorden_qa / "manifests" / "gorden-generation.json"
        generation = load_json(generation_path)
        generation["qa_results"] = generation["qa_results"][:2]
        write_json(generation_path, generation)
        expect_contract_error(
            lambda: validate_run(incomplete_gorden_qa, "preview"),
            "requires all three Gorden QA results",
            "reconstruction must retain all three Gorden QA results",
        )

        handoff_root = base / "scene1-handoff-reconstruction"
        handoff_root.mkdir()
        make_fixture(handoff_root, "reconstruction")
        upstream_root = base / "new-deck"
        attach_scene1_handoff(handoff_root, upstream_root)
        validate_run(handoff_root, "plan")
        validate_run(handoff_root, "preview")
        validate_run(handoff_root, "final")

        (handoff_root / "approvals" / "start-reconstruction.json").unlink()
        expect_contract_error(
            lambda: validate_run(handoff_root, "preview"),
            "missing user approval: start-reconstruction",
            "scene 1 handoff may not enter reconstruction without a fresh user confirmation",
        )

        missing_handoff_role = base / "missing-handoff-role"
        missing_handoff_role.mkdir()
        make_fixture(missing_handoff_role, "reconstruction")
        attach_scene1_handoff(missing_handoff_role, upstream_root)
        project_path = missing_handoff_role / "project.json"
        project = load_json(project_path)
        for source_item in project["source_files"]:
            source_item["role"] = "primary"
        write_json(project_path, project)
        expect_contract_error(
            lambda: validate_run(missing_handoff_role, "plan"),
            "requires both plans/scene1_handoff.json",
            "scene 1 handoff may not bypass validation by omitting the source role",
        )

        wrong_handoff_run = base / "wrong-handoff-run"
        wrong_handoff_run.mkdir()
        make_fixture(wrong_handoff_run, "reconstruction")
        handoff_path = attach_scene1_handoff(wrong_handoff_run, upstream_root)
        handoff = load_json(handoff_path)
        handoff["upstream_run_id"] = "not-the-upstream-run"
        write_json(handoff_path, handoff)
        expect_contract_error(
            lambda: validate_run(wrong_handoff_run, "plan"),
            "upstream project/run identity is invalid",
            "scene 1 handoff must bind the upstream run identity",
        )

        foreign_handoff_page = base / "foreign-handoff-page"
        foreign_handoff_page.mkdir()
        make_fixture(foreign_handoff_page, "reconstruction")
        handoff_path = attach_scene1_handoff(foreign_handoff_page, upstream_root)
        foreign_page = foreign_handoff_page / "foreign.png"
        write_test_png(foreign_page, 1920, 1080, 222)
        handoff = load_json(handoff_path)
        handoff["pages"][0]["path"] = str(foreign_page.resolve())
        handoff["pages"][0]["sha256"] = sha256(foreign_page)
        write_json(handoff_path, handoff)
        project_path = foreign_handoff_page / "project.json"
        project = load_json(project_path)
        project["source_files"] = [
            {
                "path": str(foreign_page.resolve()),
                "sha256": sha256(foreign_page),
                "role": "scene1-handoff",
            }
        ]
        write_json(project_path, project)
        expect_contract_error(
            lambda: validate_run(foreign_handoff_page, "plan"),
            "must be the ordered PNGs embedded in the approved image deck",
            "scene 1 handoff may not substitute unrelated PNG files",
        )

        same_run_handoff = base / "same-run-handoff"
        same_run_handoff.mkdir()
        make_fixture(same_run_handoff, "reconstruction")
        handoff_path = attach_scene1_handoff(same_run_handoff, upstream_root)
        handoff = load_json(handoff_path)
        current_project_path = same_run_handoff / "project.json"
        handoff["upstream_project"] = {
            "path": str(current_project_path.resolve()),
            "sha256": sha256(current_project_path),
        }
        write_json(handoff_path, handoff)
        expect_contract_error(
            lambda: validate_run(same_run_handoff, "plan"),
            "separate upstream run directory",
            "scene 1 and scene 3 may not share one run directory",
        )

        delegated = base / "new-deck" / DOCS["slide_plan"]
        value = load_json(delegated)
        value["status"] = "confirmed_by_delegation"
        write_json(delegated, value)
        try:
            validate_run(base / "new-deck", "preview")
        except ContractError:
            pass
        else:
            raise AssertionError("delegated confirmation must fail")

        invalid_format = base / "invalid-reconstruction-format"
        invalid_format.mkdir()
        make_fixture(invalid_format, "reconstruction")
        asset_path = invalid_format / DOCS["asset_plan"]
        value = load_json(asset_path)
        value["items"][0]["asset_type"] = "raster"
        value["items"][0]["source"] = "assets/frame-p01.jpg"
        write_json(asset_path, value)
        expect_contract_error(
            lambda: validate_run(invalid_format, "plan"),
            "asset_type must be svg or png",
            "reconstruction non-PNG raster asset must fail",
        )

        ordinary_text = base / "ordinary-text-in-png"
        ordinary_text.mkdir()
        make_fixture(ordinary_text, "reconstruction")
        asset_path = ordinary_text / DOCS["asset_plan"]
        value = load_json(asset_path)
        value["items"][0]["ordinary_text_in_asset"] = True
        write_json(asset_path, value)
        expect_contract_error(
            lambda: validate_run(ordinary_text, "plan"),
            "may not contain ordinary text",
            "ordinary text baked into reconstruction PNG must fail",
        )

        missing_png_approval = base / "missing-png-approval"
        missing_png_approval.mkdir()
        make_fixture(missing_png_approval, "reconstruction")
        (missing_png_approval / "approvals" / "reconstruction-png-assets.json").unlink()
        expect_contract_error(
            lambda: validate_run(missing_png_approval, "preview"),
            "missing user approval: reconstruction-png-assets",
            "reconstruction PNG without specific user approval must fail",
        )

        missing_gorden_scope = base / "missing-gorden-scope"
        missing_gorden_scope.mkdir()
        make_fixture(missing_gorden_scope, "reconstruction")
        (missing_gorden_scope / "approvals" / "gorden-generation-scope.json").unlink()
        expect_contract_error(
            lambda: validate_run(missing_gorden_scope, "preview"),
            "missing user approval: gorden-generation-scope",
            "Gorden generation without user scope approval must fail",
        )

        wrong_gorden_commit = base / "wrong-gorden-commit"
        wrong_gorden_commit.mkdir()
        make_fixture(wrong_gorden_commit, "reconstruction")
        component_path = wrong_gorden_commit / DOCS["gorden_component"]
        value = load_json(component_path)
        value["expected_commit"] = "0" * 40
        write_json(component_path, value)
        expect_contract_error(
            lambda: validate_run(wrong_gorden_commit, "plan"),
            "does not match expected_commit",
            "unpinned Gorden checkout must fail",
        )

        wrong_gorden_order = base / "wrong-gorden-order"
        wrong_gorden_order.mkdir()
        make_fixture(wrong_gorden_order, "new-deck")
        component_path = wrong_gorden_order / DOCS["gorden_component"]
        value = load_json(component_path)
        value["pipeline"] = "image2pptx"
        value["stages"] = [
            {
                "component": "GordenImage2PPTX",
                "purpose": "incorrectly auto-convert the scene 1 output",
                "output_root": "work/gorden/editable",
            }
        ]
        write_json(component_path, value)
        expect_contract_error(
            lambda: validate_run(wrong_gorden_order, "plan"),
            "new-deck gorden pipeline is invalid",
            "scene 1 may not auto-enter the reconstruction pipeline",
        )

        unbundled_gorden = base / "unbundled-gorden"
        unbundled_gorden.mkdir()
        make_fixture(unbundled_gorden, "reconstruction")
        component_path = unbundled_gorden / DOCS["gorden_component"]
        value = load_json(component_path)
        value["source_bundled"] = False
        write_json(component_path, value)
        expect_contract_error(
            lambda: validate_run(unbundled_gorden, "plan"),
            "must remain bundled with attribution",
            "Gorden component may not revert to an external reference contract",
        )

        understated_gorden_cost = base / "understated-gorden-cost"
        understated_gorden_cost.mkdir()
        make_fixture(understated_gorden_cost, "new-deck")
        component_path = understated_gorden_cost / DOCS["gorden_component"]
        value = load_json(component_path)
        value["generation_scope"]["estimated_imagegen_calls_min"] = 2
        write_json(component_path, value)
        expect_contract_error(
            lambda: validate_run(understated_gorden_cost, "plan"),
            "underestimates minimum imagegen calls",
            "understated Gorden imagegen scope must fail",
        )

        missing_image_deck_approval = base / "missing-image-deck-approval"
        missing_image_deck_approval.mkdir()
        make_fixture(missing_image_deck_approval, "new-deck")
        (missing_image_deck_approval / "approvals" / "image-deck-final.json").unlink()
        expect_contract_error(
            lambda: validate_run(missing_image_deck_approval, "final"),
            "missing user approval: image-deck-final",
            "scene 1 image deck without final user approval must fail",
        )

        editable_scene1 = base / "editable-scene1-final"
        editable_scene1.mkdir()
        make_fixture(editable_scene1, "new-deck")
        build_path = editable_scene1 / DOCS["build_plan"]
        value = load_json(build_path)
        value["slides"][0]["objects"] = [
            {
                "id": "wrong-native-text",
                "kind": "native_text",
                "bbox": [0, 0, 1, 1],
                "build_method": "native_text",
                "editable": True,
            }
        ]
        write_json(build_path, value)
        manifest_path = editable_scene1 / DOCS["artifacts"]
        manifest = load_json(manifest_path)
        for artifact in manifest["artifacts"]:
            if artifact.get("purpose") == "final-output":
                for parent in artifact["parents"]:
                    if Path(parent["path"]).as_posix() == Path(DOCS["build_plan"]).as_posix():
                        parent["sha256"] = sha256(build_path)
        write_json(manifest_path, manifest)
        expect_contract_error(
            lambda: validate_run(editable_scene1, "final"),
            "non-editable full-slide image",
            "scene 1 may not claim a native editable final build",
        )

        unchecked_scene1_text = base / "unchecked-scene1-text"
        unchecked_scene1_text.mkdir()
        make_fixture(unchecked_scene1_text, "new-deck")
        artifacts = load_json(unchecked_scene1_text / DOCS["artifacts"])["artifacts"]
        final_artifact = next(item for item in artifacts if item["purpose"] == "final-output")
        qa_path = unchecked_scene1_text / final_artifact["qa_path"]
        qa = load_json(qa_path)
        qa["checks"]["semantic_text_exact"] = "fail"
        write_json(qa_path, qa)
        expect_contract_error(
            lambda: validate_run(unchecked_scene1_text, "final"),
            "exact visible text and slide-plan checks",
            "scene 1 must not pass with unchecked or incorrect visible text",
        )

        fake_scene1_pptx = base / "fake-scene1-pptx"
        fake_scene1_pptx.mkdir()
        make_fixture(fake_scene1_pptx, "new-deck")
        manifest = load_json(fake_scene1_pptx / DOCS["artifacts"])
        final_artifact = next(item for item in manifest["artifacts"] if item["purpose"] == "final-output")
        write_dummy_pptx(fake_scene1_pptx / final_artifact["path"], 1)
        refresh_final_bindings(fake_scene1_pptx)
        expect_contract_error(
            lambda: validate_run(fake_scene1_pptx, "final"),
            "cannot read scene 1 image PPTX",
            "scene 1 declaration may not substitute for a real full-slide picture object",
        )

        wrong_embedded_png = base / "wrong-embedded-png"
        wrong_embedded_png.mkdir()
        make_fixture(wrong_embedded_png, "new-deck")
        alternate_png = wrong_embedded_png / "assets" / "alternate.png"
        write_test_png(alternate_png, 1920, 1080, 255)
        manifest = load_json(wrong_embedded_png / DOCS["artifacts"])
        final_artifact = next(item for item in manifest["artifacts"] if item["purpose"] == "final-output")
        write_dummy_pptx(wrong_embedded_png / final_artifact["path"], 1, [alternate_png])
        generation_path = wrong_embedded_png / "manifests" / "gorden-generation.json"
        generation = load_json(generation_path)
        wrong_slide = load_json(wrong_embedded_png / DOCS["slide_plan"])["slides"][0]
        wrong_block = wrong_slide["content_blocks"][0]
        generation["calls"].append(
            {
                "id": "call-alternate",
                "purpose": "style-option-a",
                "workflow_phase": "style-selection",
                "page_id": "P01",
                "slide_contract_sha256": json_sha256(wrong_slide),
                "style_reference_ids": load_json(
                    wrong_embedded_png / DOCS["style_reference_plan"]
                )["selected_ids"],
                "provider": "fixture-imagegen",
                "model": "fixture-model",
                "prompt": (
                    "Generate an alternate style attempt for P01\n"
                    f"{wrong_block['text']} bbox="
                    f"{json.dumps(wrong_block['bbox'], ensure_ascii=False, separators=(',', ':'))}"
                ),
                "generated_at": load_json(
                    wrong_embedded_png / "approvals" / "gorden-generation-scope.json"
                )["confirmed_at"],
                "status": "succeeded",
                "output_path": alternate_png.relative_to(wrong_embedded_png).as_posix(),
                "output_sha256": sha256(alternate_png),
            }
        )
        write_json(generation_path, generation)
        refresh_final_bindings(wrong_embedded_png)
        expect_contract_error(
            lambda: validate_run(wrong_embedded_png, "final"),
            "does not embed its confirmed PNG",
            "scene 1 PPTX must embed the exact confirmed page PNG",
        )

        damaged_qa_png = base / "damaged-qa-png"
        damaged_qa_png.mkdir()
        make_fixture(damaged_qa_png, "new-deck")
        manifest = load_json(damaged_qa_png / DOCS["artifacts"])
        final_artifact = next(item for item in manifest["artifacts"] if item["purpose"] == "final-output")
        qa = load_json(damaged_qa_png / final_artifact["qa_path"])
        (damaged_qa_png / qa["rendered_pages"][0]).write_bytes(b"PNG")
        expect_contract_error(
            lambda: validate_run(damaged_qa_png, "final"),
            "not a PNG file",
            "visual QA must reference real rendered PNG files",
        )

        incomplete_scene1_report = base / "incomplete-scene1-report"
        incomplete_scene1_report.mkdir()
        make_fixture(incomplete_scene1_report, "new-deck")
        report_path = incomplete_scene1_report / DOCS["report"]
        report = load_json(report_path)
        del report["slides"][0]["raster_assets"][0]["format"]
        write_json(report_path, report)
        expect_contract_error(
            lambda: validate_run(incomplete_scene1_report, "final"),
            "missing fields: format",
            "scene 1 report must include the complete raster disclosure",
        )

        missing_artistic_approval = base / "missing-artistic-approval"
        missing_artistic_approval.mkdir()
        make_fixture(missing_artistic_approval, "reconstruction")
        asset_path = missing_artistic_approval / DOCS["asset_plan"]
        value = load_json(asset_path)
        value["items"][0].update(
            {
                "contains_semantic_text": True,
                "semantic_text_policy": "artistic-user-confirmation",
                "text_removal_status": "artistic-only-user-confirmation",
                "visual_role": "artistic-text",
            }
        )
        write_json(asset_path, value)
        approve(missing_artistic_approval, "asset-plan", asset_path)
        approve(missing_artistic_approval, "reconstruction-png-assets", asset_path)
        expect_contract_error(
            lambda: validate_run(missing_artistic_approval, "preview"),
            "missing user approval: artistic-text-assets",
            "artistic text asset without specific user approval must fail",
        )
        approve(missing_artistic_approval, "artistic-text-assets", asset_path)
        validate_run(missing_artistic_approval, "preview")

        low_resolution = base / "low-resolution-png"
        low_resolution.mkdir()
        make_fixture(low_resolution, "reconstruction")
        asset_path = low_resolution / DOCS["asset_plan"]
        value = load_json(asset_path)
        value["items"][0]["pixel_width"] = 200
        value["items"][0]["pixel_height"] = 100
        write_json(asset_path, value)
        write_test_png(low_resolution / value["items"][0]["source"], 200, 100)
        expect_contract_error(
            lambda: validate_plan_docs(low_resolution, validate_project_file(low_resolution), "preview"),
            "resolution is too low",
            "low-resolution reconstruction PNG must fail",
        )

        unplanned_raster = base / "unplanned-raster"
        unplanned_raster.mkdir()
        make_fixture(unplanned_raster, "reconstruction")
        build_path = unplanned_raster / DOCS["build_plan"]
        value = load_json(build_path)
        value["slides"][0]["objects"][1]["asset_id"] = "not-in-plan"
        write_json(build_path, value)
        manifest_path = unplanned_raster / DOCS["artifacts"]
        manifest = load_json(manifest_path)
        for artifact in manifest["artifacts"]:
            if artifact.get("purpose") == "final-output":
                for parent in artifact["parents"]:
                    if Path(parent["path"]).as_posix() == Path(DOCS["build_plan"]).as_posix():
                        parent["sha256"] = sha256(build_path)
        write_json(manifest_path, manifest)
        expect_contract_error(
            lambda: validate_run(unplanned_raster, "final"),
            "must reference a planned asset_id",
            "unplanned reconstruction raster must fail",
        )

        extra = base / "template-fill" / "exports" / "unregistered.pptx"
        write_dummy_pptx(extra, 1)
        try:
            validate_run(base / "template-fill", "preview")
        except ContractError:
            pass
        else:
            raise AssertionError("unregistered PPTX must fail")

        stale_root = base / "stale-parent"
        stale_root.mkdir()
        make_fixture(stale_root, "existing-deck")
        stale_change_plan = stale_root / DOCS["change_plan"]
        value = load_json(stale_change_plan)
        value["changes"][0]["after"] = "确认后被篡改"
        write_json(stale_change_plan, value)
        try:
            validate_run(stale_root, "final")
        except ContractError:
            pass
        else:
            raise AssertionError("stale approval or parent hash must fail")


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("run_dir", nargs="?", help="Run directory")
    parser.add_argument("--phase", choices=sorted(PHASES), default="plan")
    parser.add_argument("--self-test", action="store_true")
    args = parser.parse_args()
    try:
        if args.self_test:
            self_test()
            print("Self-test passed: five scenarios, Gorden integration, controlled PNG reconstruction and negative hard gates")
            return 0
        if not args.run_dir:
            parser.error("run_dir is required unless --self-test is used")
        validate_run(Path(args.run_dir), args.phase)
    except ContractError as exc:
        print(f"Validation failed: {exc}")
        return 1
    print(f"Validation passed: {args.phase}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
